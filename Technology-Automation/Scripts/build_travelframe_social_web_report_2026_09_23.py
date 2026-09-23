from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]

OUT = ROOT / "Data-Analytics" / "Reports" / "2026-09-23-TravelFrame-Social-And-Website-Report.docx"

DARK = "1F2933"
NAVY = "174A63"
PALE_BLUE = "EDF5F8"
PALE_GRAY = "F6F7F8"
BORDER = "D9DDE1"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = tc_pr.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        tc_pr.append(shade)
    shade.set(qn("w:fill"), fill)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        item = borders.find(qn(tag))
        if item is None:
            item = OxmlElement(tag)
            borders.append(item)
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "6")
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), BORDER)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = margins.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def write_cell(cell, text, bold=False, color=DARK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def table(doc, headers, rows, widths, numeric_columns=None):
    numeric_columns = set(numeric_columns or [])
    item = doc.add_table(rows=1, cols=len(headers))
    item.alignment = WD_TABLE_ALIGNMENT.CENTER
    item.autofit = False
    for index, header in enumerate(headers):
        cell = item.rows[0].cells[index]
        cell.width = Inches(widths[index])
        set_cell_shading(cell, NAVY)
        write_cell(cell, header, bold=True, color="FFFFFF", size=8.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, row in enumerate(rows):
        cells = item.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = Inches(widths[index])
            set_cell_shading(cells[index], "FFFFFF" if row_index % 2 == 0 else PALE_GRAY)
            alignment = WD_ALIGN_PARAGRAPH.CENTER if index in numeric_columns else WD_ALIGN_PARAGRAPH.LEFT
            write_cell(cells[index], value, align=alignment)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return item


def heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def paragraph(doc, text, bold_lead=None):
    item = doc.add_paragraph()
    item.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        item.add_run(bold_lead).bold = True
        item.add_run(text[len(bold_lead):])
    else:
        item.add_run(text)
    return item


def bullet(doc, text):
    item = doc.add_paragraph(style="List Bullet")
    item.paragraph_format.space_after = Pt(3)
    item.add_run(text)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.62)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.3)
styles["Normal"].font.color.rgb = RGBColor.from_string(DARK)
styles["Normal"].paragraph_format.line_spacing = 1.08
styles["Normal"].paragraph_format.space_after = Pt(6)
for name, size in (("Title", 24), ("Heading 1", 16.5), ("Heading 2", 12.5)):
    styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor.from_string("000000")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("TravelFrame | Báo cáo mạng xã hội và website | 23 tháng 09 năm 2026")
footer_run.font.name = "Aptos"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor.from_string("666666")

title = doc.add_paragraph(style="Title")
title.add_run("Báo Cáo Mạng Xã Hội Và Website TravelFrame")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(16)
subtitle_run = subtitle.add_run("Facebook Fanpage Instagram Business Và Website GA4")
subtitle_run.font.name = "Aptos"
subtitle_run.font.size = Pt(13)
subtitle_run.font.color.rgb = RGBColor.from_string("000000")

table(
    doc,
    ["Trường", "Thông tin"],
    [
        ["Ngày báo cáo", "23 tháng 09 năm 2026"],
        ["Facebook", "travelframe.vn"],
        ["Instagram", "@travelframe.vn"],
        ["Website", "travelframe.sidevn.com | GA4 Property 552867189"],
        ["Social snapshot", "Meta Graph API live: 23 tháng 09 năm 2026"],
        ["Website snapshot", "GA4 live: 24 tháng 08 đến 22 tháng 09 năm 2026"],
        ["Trạng thái", "Draft để review nội bộ"],
    ],
    [1.55, 5.5],
)

heading(doc, "Tóm Tắt Điều Hành")
paragraph(
    doc,
    "TravelFrame có tín hiệu discovery tốt trên Instagram và Facebook, đặc biệt với nội dung thực dụng "
    "giúp khách quốc tế tự tin khi ở Hà Nội. Website bắt đầu nhận được lượt xem vào tour, gallery và booking, "
    "nhưng attribution chưa đủ sạch để kết luận lượng lead hoặc booking thật."
)
paragraph(
    doc,
    "Điểm nghẽn không phải là thiếu reach. Điểm nghẽn là đường đi từ nội dung sang một hành động có thể đo: "
    "keyword hoặc DM, trang đích phù hợp, form hợp lệ và đối soát với CRM."
)

heading(doc, "Tình Trạng Dữ Liệu")
table(
    doc,
    ["Nguồn", "Trạng thái", "Cách dùng trong báo cáo"],
    [
        ["Facebook và Instagram", "Đọc live thành công ngày 23 tháng 09", "Chỉ đọc dữ liệu, không có thay đổi live"],
        ["GA4", "Đọc live đến hết ngày 22 tháng 09", "Dùng làm số liệu website hiện tại"],
        ["CRM và booking", "Chưa có evidence được đối soát", "Không tính form submit là booking hoặc doanh thu"],
    ],
    [1.55, 2.15, 3.35],
)

doc.add_page_break()
heading(doc, "Bảng Điều Hành")
table(
    doc,
    ["Kênh", "Quy mô hoặc lưu lượng", "Tín hiệu tích cực", "Điểm cần xử lý"],
    [
        ["Facebook", "921 followers; 67 bài đọc được", "316 shares; bài cao nhất 66 shares", "Chưa có comments, reactions và Page Insights đầy đủ"],
        ["Instagram", "283 followers; 85 media", "137.610 reach; 16.245 interactions", "Chỉ có 11 website clicks trong 30 ngày"],
        ["Website", "64 users; 149 sessions; 742 page views", "52 views trang booking; 77 views trang tours", "34 sessions Tag Assistant và traffic admin làm nhiễu dữ liệu"],
    ],
    [1.05, 1.75, 2.25, 2.0],
)

heading(doc, "1. Facebook Fanpage")
paragraph(
    doc,
    "Dữ liệu live ngày 23 tháng 09 cho thấy Fanpage travelframe.vn có 921 followers, "
    "67 bài đọc được và 316 shares trong tập bài trả về. Shares tập trung vào một số winner, vì vậy "
    "không nên coi mọi format đều hiệu quả như nhau."
)
heading(doc, "Nội Dung Facebook Đang Thắng", level=2)
table(
    doc,
    ["Hạng", "Chủ đề", "Shares", "Lý do thắng"],
    [
        ["1", "Cách sang đường ở Việt Nam", "66", "Giải quyết nỗi lo cụ thể, dễ gửi cho người đi cùng"],
        ["2", "Hanoi Train Street quick guide", "44", "Thông tin thực tế có giá trị lên lịch và chia sẻ"],
        ["3", "5 Train Street cafés", "31", "Địa điểm cụ thể kèm bối cảnh trải nghiệm"],
        ["4", "Hidden Hanoi location kèm CTA DM", "28", "Tò mò cộng lời mời hành động rõ"],
        ["5", "Vietnam Cheat Sheet Part 2", "28", "Checklist có giá trị save và share"],
    ],
    [0.55, 2.65, 0.75, 3.1],
    numeric_columns=[0, 2],
)
paragraph(
    doc,
    "Kết luận: Facebook nên tiếp tục phân phối các series practical Hanoi utility như safety, cash, "
    "weather and packing, food routes, cheat sheets và tourist mistakes. CTA phải nối về keyword hoặc DM, "
    "không chỉ dừng ở lời mời save/share."
)

doc.add_page_break()
heading(doc, "2. Instagram Business")
paragraph(
    doc,
    "Cửa sổ Insights live là 24 tháng 08 đến 22 tháng 09 năm 2026. "
    "Instagram là động cơ discovery chính, nhưng bước từ profile sang website hiện rất thấp."
)
table(
    doc,
    ["Chỉ số", "Giá trị", "Diễn giải"],
    [
        ["Views", "330.802", "Độ phủ nội dung lớn so với quy mô follower"],
        ["Reach", "137.610", "Nhiều người mới có cơ hội nhìn thấy nội dung"],
        ["Accounts engaged", "7.519", "Có nền tảng tương tác thực"],
        ["Total interactions", "16.245", "Tương tác nội dung cao"],
        ["Profile views", "2.256", "Có cầu nối từ Reel sang profile"],
        ["Website clicks", "11", "Bottleneck lớn nhất của phễu"],
    ],
    [1.65, 1.15, 4.25],
    numeric_columns=[1],
)
heading(doc, "Nội Dung Instagram Đang Thắng", level=2)
table(
    doc,
    ["Hạng", "Nhóm nội dung", "Engagement", "Vai trò phù hợp"],
    [
        ["1", "Crossing the road in Vietnam", "345", "Safety utility và confidence building"],
        ["2", "Hanoi flooding reaction", "309", "Topical humour và discovery"],
        ["3", "Short humorous Hanoi reaction", "292", "Entertainment để duy trì reach"],
        ["4", "Hidden Hanoi location", "258", "Curiosity, local discovery và CTA DM"],
        ["5", "Save-time Vietnam tip", "240", "Utility dễ save/share"],
        ["6", "Cash-payment scam warning", "183", "Safety content có ý định cao"],
    ],
    [0.55, 2.9, 1.1, 2.5],
    numeric_columns=[0, 2],
)
paragraph(
    doc,
    "Tỷ lệ reach sang profile khoảng 1,64%, nhưng profile sang website click chỉ khoảng 0,49%. "
    "Nội dung đang tạo sự chú ý; bio, offer và CTA chưa biến sự chú ý đó thành lượt đi sâu đủ rõ."
)

doc.add_page_break()
heading(doc, "3. Website GA4")
paragraph(
    doc,
    "GA4 được truy vấn live cho 24 tháng 08 đến 22 tháng 09 năm 2026. Dữ liệu chỉ bắt đầu có hoạt động "
    "từ 14 tháng 09; vì thế không đủ lịch sử để kết luận tăng trưởng tháng hoặc so sánh tuần có ý nghĩa."
)
table(
    doc,
    ["Chỉ số", "Giá trị", "Ghi chú"],
    [
        ["Active users", "64", "Toàn bộ người dùng được GA4 ghi nhận"],
        ["Sessions", "149", "Bao gồm traffic test và nội bộ"],
        ["Page views", "742", "Có spike ngày 21 tháng 09 cần kiểm tra"],
        ["Bounce rate", "53,7%", "Không diễn giải như KPI cuối cùng khi còn traffic test"],
        ["Average session duration", "8 phút 05 giây", "Tăng bởi một số phiên sâu hoặc nội bộ"],
        ["Form starts", "20 từ 4 users", "Chưa xác nhận là khách thật"],
        ["Form submits", "19 từ 4 users", "Không được coi là booking trước khi đối soát CRM"],
    ],
    [1.7, 1.35, 3.95],
    numeric_columns=[1],
)
heading(doc, "Trang Được Xem Nhiều", level=2)
table(
    doc,
    ["Trang", "Views", "Users", "Ý nghĩa"],
    [
        ["/", "154", "59", "Trang vào chính"],
        ["/tours", "77", "10", "Bắt đầu khám phá offer"],
        ["/news", "65", "9", "Nội dung có thể hỗ trợ organic search"],
        ["/booking", "52", "11", "Có tín hiệu intent, cần đối soát"],
        ["/gallery", "29", "10", "Proof visual có người xem"],
        ["Trang quản trị", "99", "không áp dụng", "Không dùng đánh giá hiệu quả khách hàng"],
    ],
    [1.8, 0.85, 0.85, 3.5],
    numeric_columns=[1, 2],
)
heading(doc, "Nguồn Traffic", level=2)
table(
    doc,
    ["Nguồn", "Sessions", "Đọc đúng dữ liệu"],
    [
        ["Direct", "91", "Cần gắn UTM để biết nguồn thật"],
        ["tagassistant.google.com / referral", "34", "Traffic kiểm thử, cần loại khỏi KPI marketing"],
        ["google / organic", "13", "Tín hiệu SEO ban đầu"],
        ["Organic Social tổng", "10", "Chưa tương xứng với social reach"],
    ],
    [2.5, 1.0, 3.5],
    numeric_columns=[1],
)
paragraph(
    doc,
    "Ngày 21 tháng 09 có 341 page views từ 11 users. Đây là mức bất thường so với các ngày khác "
    "và cần kiểm tra log nội bộ, QA hoặc hoạt động cập nhật nội dung trước khi dùng làm tín hiệu demand."
)

heading(doc, "4. Chẩn Đoán Phễu Liên Kênh")
table(
    doc,
    ["Bước", "Hiện trạng", "Điều cần làm"],
    [
        ["Discovery", "Instagram reach và Facebook shares tốt", "Nhân bản utility, safety và local discovery"],
        ["Consideration", "Profile views, tours, gallery và booking đều có lượt xem", "Dẫn mọi content pillar đến một next step rõ"],
        ["Intent", "Website có booking views và form submits nhưng chưa sạch", "Đối soát CRM, lọc internal traffic"],
        ["Conversion", "Chưa có lead, quote, deposit hoặc paid booking được gắn nguồn", "Bắt buộc ghi source, keyword và trạng thái lead"],
    ],
    [1.25, 2.8, 2.95],
)

heading(doc, "5. Việc Ưu Tiên Trong 7 Ngày")
bullet(doc, "Lưu lịch làm mới token Meta theo scope read-only trước khi token hết hạn để duy trì snapshot social hằng tuần.")
bullet(doc, "Chốt một CTA theo ý định: HANOI cho guide/lịch trình, PHOTO cho Solo Female Photo Walk.")
bullet(doc, "Dùng UTM riêng cho Instagram, Facebook và từng content pillar; không dùng link trần chung.")
bullet(doc, "Tạo landing page cho Solo Female Photo Walk với proof, thông tin an toàn, thời lượng, CTA DM và booking.")
bullet(doc, "Loại traffic nội bộ, trang admin và Tag Assistant khỏi báo cáo GA4; đánh dấu ngày QA trong log.")
bullet(doc, "Đối soát form submit với CRM và bổ sung trường source, keyword, qualified lead, quote, deposit, paid booking.")

heading(doc, "Giới Hạn Dữ Liệu")
bullet(doc, "Social data được đọc live ngày 23 tháng 09, còn Instagram Insights dùng cửa sổ 24 tháng 08 đến 22 tháng 09.")
bullet(doc, "Facebook chưa trả comments, reactions hoặc Page Insights đầy đủ vì thiếu pages_read_user_content hoặc Page Public Content Access.")
bullet(doc, "GA4 còn traffic internal/test, đặc biệt Tag Assistant và trang quản trị.")
bullet(doc, "Không có Meta Ads spend, CRM evidence hoặc booking revenue; không tính được ROAS, CAC, CPL hoặc conversion to booking.")

heading(doc, "Kết Luận")
paragraph(
    doc,
    "TravelFrame đã chứng minh nội dung practical Hanoi utility có khả năng được khám phá, tương tác và chia sẻ. "
    "Bước tiếp theo là vận hành conversion measurement nghiêm túc: mỗi content winner cần có CTA phù hợp, mọi link "
    "cần UTM, và mọi lead phải được ghi nguồn đến khi có booking. Khi đó TravelFrame mới biết nội dung nào thực sự tạo doanh thu."
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
