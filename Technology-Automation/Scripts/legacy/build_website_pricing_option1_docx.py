from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
OUTPUT = (
    ROOT
    / "Brand-Strategy"
    / "Offers-Pricing"
    / "TravelFrame-Web-Pricing-Option-1-Implementation-Brief-2026-09-15.docx"
)

TEAL = "008AA3"
DARK = "18211D"
LIGHT = "EDF1EF"
PALE = "F5F8F7"
BORDER = "D9D9D9"
WHITE = "FFFFFF"
GRAY = "666666"
BLACK = "000000"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_col_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"

    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, DARK)
        set_cell_margins(cell)
        set_cell_border(cell)
        if widths:
            set_col_width(cell, widths[idx])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(text))
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(WHITE)
        r.font.size = Pt(font_size)

    for row_index, values in enumerate(rows):
        row = table.add_row()
        fill = WHITE if row_index % 2 == 0 else PALE
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_shading(cell, fill)
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_col_width(cell, widths[idx])
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER
                if idx > 0 and len(str(value)) < 26
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            r.font.color.rgb = RGBColor.from_string(BLACK)

    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.16
    if bold_lead:
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.12
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.32)
        p.paragraph_format.first_line_indent = Inches(-0.20)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.12
        p.add_run(f"{index}. ").bold = True
        p.add_run(item)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_code_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text)
    r.font.name = "Courier New"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor.from_string(DARK)
    return p


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(10)
    title_p_pr = title.element.get_or_add_pPr()
    title_borders = title_p_pr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_p_pr.remove(title_borders)

    for name, size in (("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)

    for style_name in ("List Bullet", "List Number"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.size = Pt(10.5)
        styles[style_name].font.color.rgb = RGBColor.from_string(BLACK)

    header = section.header.paragraphs[0]
    header.text = "TravelFrame Vietnam   Web implementation brief"
    header.style = styles["Normal"]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

    footer = section.footer.paragraphs[0]
    footer.add_run("Internal working document   2026 09 15")
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    footer.add_run("     ")
    add_page_number(footer)


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("TravelFrame Website Pricing Option 1 Implementation Brief")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    r = subtitle.add_run(
        "Hướng dẫn cho team web chuyển sang mô hình tour trọn gói đã bao gồm ảnh và video"
    )
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(DARK)

    metadata = add_table(
        doc,
        ["Thông tin", "Nội dung"],
        [
            ("Ngày lập", "2026-09-15"),
            ("Người đọc", "Team web, content, vận hành và người duyệt giá"),
            ("Website", "travelframe.sidevn.com"),
            ("Quyết định", "Chọn Phương án 1: giá tour đã bao gồm một gói ảnh hoặc video mặc định"),
            ("Trạng thái", "Implementation brief. Giá thương mại cần owner xác nhận trước khi publish"),
        ],
        widths=[1.55, 5.15],
        font_size=10,
    )
    metadata.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_body(
        doc,
        "Tài liệu này giải thích cách sửa website để khách nhìn thấy một mức giá tour dễ hiểu. "
        "Giá trên tour card sẽ là giá của một trải nghiệm hoàn chỉnh, đã gồm một mức ảnh hoặc video mặc định. "
        "Khách chỉ trả thêm khi chủ động nâng cấp gói hình ảnh hoặc chọn dịch vụ phát sinh.",
    )
    add_body(
        doc,
        "Kết quả cần đạt là khách không còn gặp hai cách tính giá khác nhau giữa trang tour và trang booking. "
        "Team web có thể thay đổi giá, gói mặc định và phụ phí trong dữ liệu quản trị mà không phải sửa code.",
    )

    doc.add_page_break()

    add_heading(doc, "1 Mục tiêu và phạm vi")
    add_heading(doc, "1.1 Mục tiêu", level=2)
    add_bullets(
        doc,
        [
            "Mỗi tour có một giá công bố dễ hiểu và đã bao gồm ảnh hoặc video.",
            "Visual package trở thành lựa chọn nâng cấp, không còn là khoản bắt buộc cộng thêm.",
            "Giá trên tour card, tour detail, booking summary và confirmation phải giống nhau.",
            "Duration, group size, delivery time và inclusions chỉ có một nguồn dữ liệu.",
            "Các giá chưa được duyệt không được xuất hiện trên website live.",
        ],
    )

    add_heading(doc, "1.2 Phạm vi cần sửa", level=2)
    add_bullets(
        doc,
        [
            "Homepage.",
            "Tours catalogue.",
            "Tour detail pages.",
            "Booking Step 2 Packages.",
            "Booking calculator và Review and book.",
            "Booking confirmation.",
            "FAQ hoặc How it works nếu đang nhắc tới giá và delivery.",
            "Dữ liệu quản trị cho tour, visual package và upgrade surcharge.",
        ],
    )

    add_heading(doc, "1.3 Không nằm trong lần sửa này", level=2)
    add_bullets(
        doc,
        [
            "Không thay đổi brand architecture Hanoi Expert, TravelFrame và The Hanoi Edit.",
            "Không tự quyết định giá thương mại cuối cùng.",
            "Không thay đổi legal model hoặc payment gateway.",
            "Không thêm tour mới.",
            "Không tự động publish khi chưa có người duyệt giá và nội dung.",
        ],
    )

    add_heading(doc, "2 Vấn đề hiện tại")
    add_table(
        doc,
        ["Vấn đề", "Hiện tại", "Trạng thái mục tiêu"],
        [
            (
                "Ảnh và video",
                "Tour card ghi đã bao gồm nhưng booking bắt khách mua thêm",
                "Mỗi tour có visual tier mặc định đã nằm trong giá",
            ),
            (
                "Đơn vị giá",
                "Tour card ghi per person nhưng calculator không nhân số khách",
                "Mỗi tour khai báo rõ per group hoặc per person và calculator làm đúng",
            ),
            (
                "Delivery",
                "Homepage ghi 24 đến 48 giờ nhưng package ghi 1 giờ, 2 ngày hoặc 3 ngày",
                "Hiển thị delivery theo visual tier được chọn",
            ),
            (
                "Duration",
                "Food Tour và Cinematic Hanoi có hai duration khác nhau",
                "Mọi trang đọc cùng một duration field",
            ),
            (
                "Group size",
                "Tour và package có capacity không khớp",
                "Booking dùng giới hạn nhỏ hơn giữa route và package",
            ),
            (
                "Giá cuối",
                "Khách khó biết tour card có phải số tiền cuối hay không",
                "Tour card là giá trọn gói cơ bản, add on ghi riêng",
            ),
        ],
        widths=[1.35, 2.65, 2.70],
    )

    add_heading(doc, "3 Quy tắc sản phẩm mới")
    add_heading(doc, "3.1 Quy tắc chính", level=2)
    add_numbered(
        doc,
        [
            "Mỗi tour phải có một included visual tier.",
            "Giá tour cơ bản đã bao gồm visual tier đó.",
            "Khách không cần chọn lại gói đang được bao gồm.",
            "Khách có thể giữ gói mặc định hoặc chọn gói cao hơn.",
            "Website chỉ cộng upgrade surcharge và các add on thật sự phát sinh.",
            "Final total phải giống nhau ở summary, review, booking payload và confirmation.",
        ],
    )

    add_heading(doc, "3.2 Công thức giá", level=2)
    add_code_line(
        doc,
        "Estimated total = published tour price + visual upgrade surcharge + paid add ons",
    )
    add_body(
        doc,
        "Không tiếp tục dùng công thức Visual package price nhân với toàn bộ duration của tour trong customer booking flow. "
        "Mỗi mức nâng cấp nên có phụ phí rõ ràng theo từng tour. Cách này giúp team business kiểm soát margin và tránh việc tour dài tạo ra tổng giá bất ngờ.",
    )

    add_heading(doc, "3.3 Đơn vị giá", level=2)
    add_body(
        doc,
        "Khuyến nghị launch là giá theo private group vì TravelFrame bán private experience và chi phí guide hoặc photographer phần lớn gắn với một booking. "
        "Tuy nhiên hệ thống nên hỗ trợ cả per group và per person để không khóa cứng mô hình tương lai.",
    )
    add_table(
        doc,
        ["price unit", "Calculator", "Cách hiển thị"],
        [
            (
                "per_group",
                "Cộng tour price một lần. Có thể cộng extra guest fee nếu vượt số khách bao gồm",
                "From 6.500.000 VND per private group up to 4 guests",
            ),
            (
                "per_person",
                "Nhân giá với số khách tính phí. Áp dụng child rule nếu đã được duyệt",
                "From 1.300.000 VND per person",
            ),
        ],
        widths=[1.15, 3.05, 2.50],
    )

    add_heading(doc, "4 Cấu hình dữ liệu đề xuất")
    add_body(
        doc,
        "Các trường dưới đây nên nằm trong CMS hoặc database. Team content và business có thể sửa dữ liệu mà không sửa template hoặc JavaScript.",
    )
    add_table(
        doc,
        ["Field", "Bắt buộc", "Ý nghĩa", "Ví dụ"],
        [
            ("published_price_vnd", "Có", "Giá tour đã gồm visual mặc định", "6.500.000"),
            ("display_price_usd", "Có", "USD hiển thị sau khi business duyệt", "250"),
            ("price_unit", "Có", "per_group hoặc per_person", "per_group"),
            ("included_guest_count", "Có", "Số khách nằm trong giá group", "4"),
            ("max_guest_count", "Có", "Số khách tối đa của route", "4"),
            ("extra_guest_fee_vnd", "Tùy chọn", "Phụ phí mỗi khách vượt mức bao gồm", "500.000"),
            ("included_visual_package_id", "Có", "Visual tier mặc định của tour", "explorer"),
            ("visual_upgrade_options", "Có", "Các tier được phép nâng cấp", "complete vip"),
            ("upgrade_surcharge_vnd", "Có", "Phụ phí cố định theo tour và tier", "2.000.000"),
            ("duration_minutes", "Có", "Nguồn duration duy nhất", "360"),
            ("delivery_hours", "Có", "Thời gian giao theo tier", "72"),
            ("price_status", "Có", "draft approved hoặc hidden", "approved"),
            ("last_price_reviewed_at", "Có", "Ngày duyệt giá gần nhất", "2026-09-15"),
        ],
        widths=[1.65, 0.72, 2.70, 1.63],
        font_size=8.8,
    )

    add_heading(doc, "5 Gói hình ảnh")
    add_body(
        doc,
        "Tên và deliverables có thể giữ như hiện tại. Trong tour flow, không hiển thị giá theo giờ cho tier đã được bao gồm. "
        "Nếu TravelFrame vẫn bán photoshoot rời, giá theo giờ có thể nằm ở một trang dịch vụ riêng.",
    )
    add_table(
        doc,
        ["Tier", "Deliverable hiện tại", "Cách dùng trong tour flow"],
        [
            ("Photo Phone", "Ảnh và video bằng iPhone, giao trong 1 giờ", "Included tier hoặc lựa chọn cơ bản"),
            ("Explorer", "10 ảnh máy chuyên nghiệp, giao trong 2 ngày", "Included tier phổ biến"),
            ("Complete", "25 ảnh retouch và video dưới 1 phút trong 3 ngày", "Included tier cho tour thiên về hình ảnh hoặc upgrade"),
            ("VIP", "50 ảnh retouch và cinematic video dưới 3 phút trong 3 ngày", "Upgrade cao nhất"),
        ],
        widths=[1.15, 3.15, 2.55],
    )

    add_heading(doc, "6 Mapping đề xuất cho bảy tour")
    add_body(
        doc,
        "Bảng này là đề xuất để team web chuẩn bị cấu hình. Owner cần duyệt lại giá và margin trước khi chuyển price status sang approved.",
    )
    add_table(
        doc,
        ["Tour", "Duration", "Visual mặc định", "Giá website hiện tại", "Việc cần chốt"],
        [
            ("Spiritual and West Lake", "6 giờ", "Explorer", "6.500.000 VND", "Kiểm tra margin xe vé và food stops"),
            ("French Colonial", "6 giờ", "Explorer", "6.500.000 VND", "Kiểm tra margin xe vé và photographer"),
            ("Hanoi Early Bird", "5 giờ", "Explorer", "4.150.000 VND", "Chốt breakfast và pickup cost"),
            ("Colors of the Delta", "10 giờ", "Complete", "6.500.000 VND", "Bắt buộc định giá lại vì có xe vé và lunch"),
            ("Hanoi Unfiltered", "5 giờ", "Explorer", "5.200.000 VND", "Train Street phải là optional stop"),
            ("A Taste of Hanoi", "3 giờ", "Photo Phone", "1.300.000 VND", "Chốt food cost và duration"),
            ("Cinematic Hanoi", "3 giờ", "Complete", "3.900.000 VND", "Bắt buộc kiểm tra margin và duration"),
        ],
        widths=[1.55, 0.72, 1.20, 1.32, 2.00],
        font_size=8.7,
    )

    add_heading(doc, "7 Thay đổi theo từng trang")
    add_heading(doc, "7.1 Homepage", level=2)
    add_bullets(
        doc,
        [
            "Giữ thông điệp Every tour includes photo and video.",
            "Bỏ câu giao toàn bộ trong 24 đến 48 giờ.",
            "Thay bằng Delivery depends on your included or upgraded visual package.",
            "Không hiển thị bốn Visual package như bốn khoản khách bắt buộc phải cộng thêm.",
            "Có thể đổi section thành Included visual coverage and optional upgrades.",
        ],
    )

    add_heading(doc, "7.2 Tours catalogue", level=2)
    add_bullets(
        doc,
        [
            "Mỗi card hiển thị tour price, price unit, guest count và included visual tier.",
            "Không dùng From nếu tour chỉ có một giá cơ bản đã xác định.",
            "Nếu vẫn dùng From, phải có lý do rõ như group size hoặc seasonal surcharge.",
            "Dòng mẫu: 6.500.000 VND per private group, includes Explorer photo package.",
        ],
    )

    add_heading(doc, "7.3 Tour detail", level=2)
    add_bullets(
        doc,
        [
            "Thêm section What your price includes.",
            "Hiển thị included visual tier và deliverables cụ thể.",
            "Hiển thị upgrade options dưới dạng phụ phí.",
            "Dùng duration_minutes làm nguồn duy nhất cho hero, overview và itinerary.",
            "Dùng max_guest_count làm nguồn duy nhất cho capacity.",
            "Không dùng daily departures nếu chưa xác nhận lịch guide.",
        ],
    )

    add_heading(doc, "7.4 Booking Step 2", level=2)
    add_bullets(
        doc,
        [
            "Khách chọn tour trước.",
            "Hệ thống tự gắn included visual tier.",
            "Card visual mặc định hiển thị Included in your tour.",
            "Chỉ các tier cao hơn mới có nút Upgrade và phụ phí.",
            "Không cho downgrade nếu điều đó làm tour không còn đúng promise trên tour card.",
            "Không cho tiếp tục nếu tour price status chưa approved.",
        ],
    )

    add_heading(doc, "7.5 Review and book", level=2)
    add_table(
        doc,
        ["Dòng hiển thị", "Ví dụ"],
        [
            ("Tour", "Hanoi Early Bird"),
            ("Tour price", "4.150.000 VND per private group"),
            ("Included visual package", "Explorer included"),
            ("Visual upgrade", "Complete plus 1.800.000 VND"),
            ("Paid add ons", "Additional stop plus 390.000 VND"),
            ("Estimated total", "6.340.000 VND"),
            ("Final note", "Final availability and price are confirmed by TravelFrame"),
        ],
        widths=[2.05, 4.65],
    )

    add_heading(doc, "7.6 Booking confirmation", level=2)
    add_bullets(
        doc,
        [
            "Lưu lại tour price, price unit và included visual package tại thời điểm booking.",
            "Không chỉ lưu ID rồi đọc giá mới nhất sau này.",
            "Confirmation phải ghi rõ số khách, delivery time, inclusions và upgrade.",
            "Nếu team thay giá sau đó, booking cũ vẫn giữ snapshot giá đã gửi.",
        ],
    )

    add_heading(doc, "8 Logic calculator")
    add_heading(doc, "8.1 Luồng đơn giản", level=2)
    add_numbered(
        doc,
        [
            "Đọc published tour price.",
            "Kiểm tra price unit.",
            "Nếu per group, cộng tour price một lần và cộng extra guest fee khi cần.",
            "Nếu per person, tính adult và child subtotal theo rule đã duyệt.",
            "Gắn included visual tier với giá nâng cấp bằng 0.",
            "Nếu khách chọn tier cao hơn, cộng upgrade surcharge của chính tour đó.",
            "Cộng paid add ons.",
            "Hiển thị cùng một kết quả tại live summary, review và confirmation.",
        ],
    )

    add_heading(doc, "8.2 Pseudocode", level=2)
    add_code_line(doc, "total = calculateTourPrice(tour, guests)")
    add_code_line(doc, "total += getVisualUpgradeSurcharge(tour, selectedVisualTier)")
    add_code_line(doc, "total += calculatePaidAddOns(selectedItems)")
    add_code_line(doc, "return total")

    add_heading(doc, "8.3 Quy tắc USD và VND", level=2)
    add_bullets(
        doc,
        [
            "VND nên là giá gốc trong database.",
            "USD chỉ hiển thị khi đã được business duyệt hoặc được quy đổi theo một exchange rate có ngày cập nhật.",
            "Không nhập hai giá độc lập mà không có bước kiểm tra.",
            "Booking confirmation phải ghi currency thực tế khách sẽ thanh toán.",
        ],
    )

    add_heading(doc, "9 Nội dung cần sửa ngay")
    add_table(
        doc,
        ["Vị trí", "Nội dung cũ", "Nội dung mới"],
        [
            ("Homepage delivery", "Color graded delivered in 24 to 48h", "Delivery follows your included or upgraded visual package"),
            ("Tour cards", "Photo and video included plus From price per person", "Includes named visual tier plus one clear price unit"),
            ("Food Tour", "3 hours ở hero nhưng 5 hours trong overview", "Dùng một duration đã được owner xác nhận"),
            ("Cinematic Hanoi", "3 hours ở hero nhưng 6 hours trong overview", "Dùng một duration đã được owner xác nhận"),
            ("Availability", "Daily departures", "Available daily subject to confirmation"),
            ("Train Street", "Điểm cố định hoặc access ngầm được đảm bảo", "Optional stop depending on local access and restrictions"),
            ("Exclusive access", "Claim chưa có bằng chứng", "Bỏ hoặc thay bằng curated local route"),
        ],
        widths=[1.35, 2.55, 2.80],
        font_size=9,
    )

    add_heading(doc, "10 Quy tắc trẻ em và số khách")
    add_body(
        doc,
        "Nếu tour dùng per group, không tự động giảm 50 phần trăm tour price cho trẻ em. "
        "Chi phí vé, đồ ăn, ghế xe hoặc dịch vụ phát sinh được tính riêng nếu cần. "
        "Nếu tour dùng per person, child rule chỉ được bật sau khi business xác nhận.",
    )
    add_table(
        doc,
        ["Tình huống", "Kết quả mong đợi"],
        [
            ("Per group, 2 người lớn và 1 trẻ em nằm trong capacity", "Tour base không đổi. Chỉ cộng child specific costs nếu có"),
            ("Per group, số khách vượt included count nhưng chưa vượt max", "Cộng extra guest fee"),
            ("Per person", "Nhân đúng số khách tính phí và áp child rule được duyệt"),
            ("Vượt max guest count", "Không cho book. Chuyển sang Customize hoặc human quote"),
        ],
        widths=[2.80, 3.90],
    )

    add_heading(doc, "11 Trạng thái giá và quyền publish")
    add_table(
        doc,
        ["price status", "Ý nghĩa", "Website live"],
        [
            ("draft", "Đang nhập hoặc đang tính margin", "Không hiển thị và không cho book"),
            ("approved", "Owner đã duyệt giá, tier và capacity", "Được hiển thị và book"),
            ("hidden", "Tạm dừng bán", "Không xuất hiện trong catalogue"),
        ],
        widths=[1.25, 3.25, 2.20],
    )
    add_body(
        doc,
        "Mỗi lần đổi giá cần lưu người duyệt, thời điểm duyệt và ghi chú. Team web không tự chuyển draft sang approved.",
    )

    add_heading(doc, "12 Test cases bắt buộc")
    add_table(
        doc,
        ["ID", "Kịch bản", "Kết quả đạt"],
        [
            ("T01", "Mở tour card và tour detail", "Giá, unit, duration và capacity giống nhau"),
            ("T02", "Chọn tour không nâng cấp visual", "Visual mặc định giá 0 và total bằng published tour price"),
            ("T03", "Nâng Explorer lên Complete", "Chỉ cộng đúng surcharge được cấu hình cho tour"),
            ("T04", "Đổi lại included tier", "Surcharge trở về 0"),
            ("T05", "Per group với nhiều khách trong included count", "Tour base không bị nhân"),
            ("T06", "Per person với 3 người lớn", "Tour base được nhân 3"),
            ("T07", "Vượt max guest count", "Không cho tiếp tục và đề nghị Customize"),
            ("T08", "Thêm một paid add on", "Giá và duration cập nhật một lần"),
            ("T09", "Review và confirmation", "Total và line items giống live summary"),
            ("T10", "Đổi giá sau một booking cũ", "Booking cũ giữ nguyên price snapshot"),
            ("T11", "Tour price status là draft", "Tour không thể được book"),
            ("T12", "Kiểm tra mobile", "Không cắt giá, tên package hoặc nút upgrade"),
        ],
        widths=[0.55, 3.15, 3.20],
        font_size=8.8,
    )

    add_heading(doc, "13 Kế hoạch triển khai")
    add_heading(doc, "13.1 Bước 1 Chốt dữ liệu", level=2)
    add_bullets(
        doc,
        [
            "Owner duyệt price unit cho từng tour.",
            "Owner duyệt published price và included visual tier.",
            "Owner duyệt upgrade surcharge.",
            "Content xác nhận duration, capacity, inclusions và delivery.",
        ],
    )

    add_heading(doc, "13.2 Bước 2 Sửa CMS và backend", level=2)
    add_bullets(
        doc,
        [
            "Thêm các field trong mục 4.",
            "Thêm price status và approval metadata.",
            "Thêm price snapshot vào booking.",
            "Sửa calculator theo mục 8.",
        ],
    )

    add_heading(doc, "13.3 Bước 3 Sửa giao diện", level=2)
    add_bullets(
        doc,
        [
            "Sửa Homepage, cards, tour details và booking Step 2.",
            "Hiển thị included tier trước upgrade options.",
            "Hiển thị line item rõ ràng trong review.",
            "Kiểm tra desktop và mobile.",
        ],
    )

    add_heading(doc, "13.4 Bước 4 QA và publish", level=2)
    add_bullets(
        doc,
        [
            "Chạy đủ test cases.",
            "Content review toàn bộ English và Vietnamese.",
            "Business kiểm tra ba booking mẫu có giá thấp, trung bình và cao.",
            "Chỉ publish khi mọi tour live có price status approved.",
            "Giữ bản backup để rollback nếu booking total sai.",
        ],
    )

    add_heading(doc, "14 Checklist duyệt trước khi giao team web")
    add_table(
        doc,
        ["Duyệt", "Nội dung", "Owner"],
        [
            ("☐", "Chọn per group hay per person cho từng tour", "Business"),
            ("☐", "Duyệt giá tour trọn gói", "Business"),
            ("☐", "Duyệt included visual tier", "Business và Production"),
            ("☐", "Duyệt upgrade surcharge", "Business và Production"),
            ("☐", "Duyệt duration và group size", "Operations"),
            ("☐", "Duyệt delivery time", "Production"),
            ("☐", "Duyệt inclusions và exclusions", "Operations"),
            ("☐", "Duyệt nội dung Train Street và availability", "Operations"),
            ("☐", "Duyệt English và Vietnamese copy", "Content"),
        ],
        widths=[0.65, 4.45, 1.60],
    )

    add_heading(doc, "15 Điều kiện hoàn thành")
    add_bullets(
        doc,
        [
            "Không còn trường hợp tour card nói included nhưng checkout cộng lại toàn bộ Visual package.",
            "Khách hiểu giá đang tính theo group hay person trước khi bấm Book.",
            "Mọi trang hiển thị cùng một duration, capacity và delivery rule.",
            "Total không thay đổi giữa live summary, review và confirmation.",
            "Team business có thể sửa giá và surcharge trong CMS.",
            "Tour chưa duyệt không thể được publish hoặc book.",
            "Không còn claim 24 đến 48 giờ áp dụng cho mọi package.",
        ],
    )

    add_heading(doc, "16 Tài liệu nguồn")
    add_bullets(
        doc,
        [
            "Website travelframe.sidevn.com được kiểm tra ngày 2026-09-15.",
            "TravelFrame Website Packages and Pricing snapshot ngày 2026-09-15.",
            "TravelFrame Decision Log.",
            "TravelFrame Product and Guest Experience.",
            "TravelFrame Operations and Policies.",
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
