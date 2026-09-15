from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "docs/superpowers/specs/2026-09-14-travelframe-30-day-content-design.md"
OUTPUT = (
    ROOT
    / "Marketing-Growth"
    / "Plans"
    / "TravelFrame-30-Day-Content-Plan-2026-09-14.docx"
)

BLACK = "000000"
CHARCOAL = "333333"
MID_GRAY = "666666"
LIGHT_GRAY = "D9D9D9"
PALE_GRAY = "F4F6F7"
TEAL = "008AA3"
PALE_TEAL = "E8F5F7"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=LIGHT_GRAY, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_fixed_cell_width(cell, inches: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_rich_text(paragraph, text: str, color=CHARCOAL, size=10.5) -> None:
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        is_code = part.startswith("`") and part.endswith("`")
        clean = part[2:-2] if is_bold else part[1:-1] if is_code else part
        run = paragraph.add_run(clean)
        run.bold = is_bold
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        if is_code:
            run.font.name = "Aptos Mono"


def style_paragraph(paragraph, before=0, after=5, line=1.12) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_body_paragraph(doc, text: str, after=6):
    p = doc.add_paragraph()
    add_rich_text(p, text, size=10.8)
    style_paragraph(p, after=after, line=1.16)
    return p


def add_bullet(doc, text: str, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    add_rich_text(p, text, size=10.5)
    style_paragraph(p, after=2, line=1.1)
    return p


def add_numbered(doc, text: str, number: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.34)
    p.paragraph_format.first_line_indent = Inches(-0.24)
    number_run = p.add_run(f"{number}.")
    number_run.font.name = "Aptos"
    number_run.font.size = Pt(10.5)
    number_run.font.color.rgb = RGBColor.from_string(CHARCOAL)
    p.add_run("\t")
    add_rich_text(p, text, size=10.5)
    style_paragraph(p, after=2, line=1.1)
    return p


def add_heading(doc, text: str, level: int):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_rich_text(p, re.sub(r"^\d+\.\s*", "", text), color=BLACK, size={1: 17, 2: 13, 3: 11.5}.get(level, 11))
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    return p


def style_table(table, widths=None) -> None:
    table.autofit = False
    for row_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        if row_idx == 0:
            set_repeat_table_header(row)
        for col_idx, cell in enumerate(row.cells):
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_fixed_cell_width(cell, widths[col_idx])
            set_cell_shading(cell, TEAL if row_idx == 0 else (PALE_GRAY if row_idx % 2 == 0 else WHITE))
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor.from_string(WHITE if row_idx == 0 else CHARCOAL)
                    if row_idx == 0:
                        run.bold = True


def add_markdown_table(doc, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    headers = rows[0]
    body = [row for row in rows[2:] if row]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header.strip()
    for row in body:
        cells = table.add_row().cells
        for idx, value in enumerate(row[: len(headers)]):
            cells[idx].text = re.sub(r"\*\*|`", "", value.strip())
    page_width = 7.0
    if len(headers) == 3:
        widths = [4.3, 1.25, 1.45]
    elif len(headers) == 2:
        widths = [2.2, 4.8]
    else:
        widths = [page_width / len(headers)] * len(headers)
    style_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def render_generic_markdown(doc, lines: list[str]) -> None:
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:], 1)
        elif line.startswith("### "):
            add_heading(doc, line[4:], 2)
        elif line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            add_markdown_table(doc, table_lines)
            continue
        elif match := re.match(r"^(\d+)\.\s+(.+)", line):
            add_numbered(doc, match.group(2), match.group(1))
        elif line.startswith("- "):
            add_bullet(doc, line[2:])
        elif line.startswith("# "):
            pass
        else:
            add_body_paragraph(doc, line)
        i += 1


def parse_calendar(lines: list[str]) -> list[dict[str, str]]:
    entries = []
    for line in lines:
        if not line.startswith("| 2026-"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) != 5:
            continue
        entries.append(
            {
                "date": cells[0],
                "pillar": cells[1],
                "reel": cells[2],
                "story": cells[3],
                "carousel": cells[4],
            }
        )
    return entries


def add_day_record(doc, item: dict[str, str]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.style = "Table Grid"
    row = table.rows[0]
    prevent_row_split(row)
    left, right = row.cells
    set_fixed_cell_width(left, 1.35)
    set_fixed_cell_width(right, 5.65)
    for cell in (left, right):
        set_cell_borders(cell)
        set_cell_margins(cell, top=130, start=140, bottom=130, end=140)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(left, PALE_TEAL)
    set_cell_shading(right, WHITE)

    left.text = ""
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(item["date"])
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(11.5)
    run.font.color.rgb = RGBColor.from_string(BLACK)
    style_paragraph(p, after=5, line=1.0)
    p = left.add_paragraph()
    add_rich_text(p, item["pillar"], color=TEAL, size=9.5)
    style_paragraph(p, after=0, line=1.0)

    right.text = ""
    for idx, (label, text) in enumerate(
        (
            ("Reel", item["reel"]),
            ("Story", item["story"]),
            ("Carousel", item["carousel"]),
        )
    ):
        p = right.paragraphs[0] if idx == 0 else right.add_paragraph()
        label_run = p.add_run(f"{label}: ")
        label_run.bold = True
        label_run.font.name = "Aptos"
        label_run.font.size = Pt(10)
        label_run.font.color.rgb = RGBColor.from_string(BLACK)
        add_rich_text(p, text, size=10)
        style_paragraph(p, after=3 if idx < 2 else 0, line=1.08)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(CHARCOAL)

    for name, size in (("Title", 27), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor.from_string(BLACK)
        if name == "Title":
            p_pr = style.element.get_or_add_pPr()
            p_bdr = p_pr.find(qn("w:pBdr"))
            if p_bdr is not None:
                p_pr.remove(p_bdr)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "TRAVELFRAME  CONTENT OPERATIONS"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("TravelFrame  |  ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_field(fp, "PAGE")


def add_cover(doc: Document) -> None:
    doc.add_paragraph().paragraph_format.space_after = Pt(72)
    eyebrow = doc.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = eyebrow.add_run("CONTENT OPERATIONS PLAN")
    run.font.name = "Aptos"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(TEAL)
    style_paragraph(eyebrow, after=14)

    title = doc.add_paragraph(style="Title")
    title.add_run("TravelFrame 30 Day Content Plan")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(16)
    remove_paragraph_borders(title)

    subtitle = doc.add_paragraph()
    add_rich_text(
        subtitle,
        "Lịch nội dung Instagram từ 2026-09-15 đến 2026-10-14",
        color=CHARCOAL,
        size=14,
    )
    style_paragraph(subtitle, after=18, line=1.1)

    intro = doc.add_paragraph()
    add_rich_text(
        intro,
        "Kế hoạch này tổ chức 30 Reel, 8 Carousel và hệ thống Story hằng ngày để TravelFrame vừa cung cấp thông tin hữu ích cho khách quốc tế, vừa xây dựng nhu cầu chụp ảnh và đặt trải nghiệm tại Hà Nội.",
        color=CHARCOAL,
        size=11.5,
    )
    intro.paragraph_format.right_indent = Inches(0.6)
    style_paragraph(intro, after=24, line=1.2)

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    values = (
        ("Thời gian", "2026-09-15 đến 2026-10-14"),
        ("Nhịp đăng", "1 Reel/ngày  |  2-4 Story/ngày  |  2 Carousel/tuần"),
        ("Primary conversion", "Instagram DM với keyword HANOI"),
    )
    for idx, (label, value) in enumerate(values):
        meta.rows[idx].cells[0].text = label
        meta.rows[idx].cells[1].text = value
        for cell in meta.rows[idx].cells:
            set_cell_borders(cell)
            set_cell_margins(cell, top=120, start=140, bottom=120, end=140)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, PALE_TEAL if idx % 2 == 0 else WHITE)
        for run in meta.rows[idx].cells[0].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(BLACK)
        for cell in meta.rows[idx].cells:
            for p in cell.paragraphs:
                style_paragraph(p, after=0, line=1.05)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(10)
    set_fixed_cell_width(meta.rows[0].cells[0], 1.65)
    set_fixed_cell_width(meta.rows[0].cells[1], 5.35)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_rich_text(date_p, "Prepared 2026-09-14", color=MID_GRAY, size=9)
    date_p.paragraph_format.space_after = Pt(0)
    doc.add_page_break()


def build() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    calendar_start = lines.index("## 5. Lịch nội dung 30 ngày")
    production_start = lines.index("## 6. Nhịp sản xuất theo tuần")

    before_calendar = lines[1:calendar_start]
    calendar_lines = lines[calendar_start:production_start]
    after_calendar = lines[production_start:]
    calendar = parse_calendar(calendar_lines)

    if len(calendar) != 30:
        raise RuntimeError(f"Expected 30 calendar entries, found {len(calendar)}")

    doc = Document()
    configure_document(doc)
    add_cover(doc)

    render_generic_markdown(doc, before_calendar)

    add_heading(doc, "Lịch nội dung 30 ngày", 1)
    add_body_paragraph(
        doc,
        "Lịch được chia thành năm chặng để đội ngũ quay theo batch, theo dõi nhịp series và tái sử dụng footage cho Story hoặc Carousel.",
        after=10,
    )
    for idx, item in enumerate(calendar):
        if idx % 7 == 0:
            week_number = idx // 7 + 1
            end_idx = min(idx + 7, len(calendar)) - 1
            add_heading(
                doc,
                f"Tuần {week_number}  {calendar[idx]['date']} đến {calendar[end_idx]['date']}",
                2,
            )
        add_day_record(doc, item)

    render_generic_markdown(doc, after_calendar)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
