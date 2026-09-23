from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]

OUT = ROOT / "Data-Analytics" / "Reports" / "2026-09-18-TravelFrame-Meta-GA-Combined-Report.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

TEAL = "008AA3"
DARK = "222222"
LIGHT_GRAY = "F4F6F7"
BORDER = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)


def cell_text(cell, value, bold=False, color=DARK, size=9.2):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(value))
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)


def make_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Inches(widths[i])
        shade(c, TEAL)
        cell_text(c, h, bold=True, color="FFFFFF", size=8.8)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            shade(cells[i], "FFFFFF" if ridx % 2 == 0 else LIGHT_GRAY)
            cell_text(cells[i], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.4)
styles["Normal"].font.color.rgb = RGBColor.from_string(DARK)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12.5)):
    styles[name].font.name = "Aptos Display" if name == "Title" else "Aptos"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor.from_string("000000")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("TravelFrame | Báo cáo Meta và GA4 | 18 tháng 09 năm 2026")
fr.font.name = "Aptos"
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor.from_string("666666")

title = doc.add_paragraph(style="Title")
title.add_run("Báo Cáo Tổng Hợp Kênh TravelFrame")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
sr = subtitle.add_run("Facebook Fanpage Instagram và Website GA4")
sr.font.name = "Aptos"
sr.font.size = Pt(14)
sr.font.color.rgb = RGBColor.from_string("000000")

make_table(
    doc,
    ["Trường", "Thông tin"],
    [
        ["Ngày báo cáo", "18 tháng 09 năm 2026"],
        ["Facebook", "travelframe.vn | Page ID 1409099405611405"],
        ["Instagram", "@travelframe.vn | Business Account ID 17841438272497512"],
        ["Website", "https://travelframe.sidevn.com | GA4 Property 552867189"],
        ["Nguồn dữ liệu", "Meta Graph API qua MCP read-only và dữ liệu GA4 do người dùng cung cấp"],
        ["Trạng thái", "Đọc live / Bản nháp để review"],
    ],
    [1.45, 5.45],
)

h(doc, "Tóm Tắt Điều Hành")
doc.add_paragraph(
    "TravelFrame đang hoạt động tốt nhất ở vai trò thương hiệu cung cấp thông tin du lịch Hà Nội và Việt Nam "
    "thực tế, dễ lưu và dễ chia sẻ. Instagram là động cơ khám phá và tạo tương tác lớn; Facebook đang có quy mô "
    "follower và tín hiệu chia sẻ tăng mạnh; GA4 cho thấy một số khách quốc tế đã đi sâu tới trang tour, thư viện ảnh "
    "và booking."
)
doc.add_paragraph(
    "Cơ hội ưu tiên không chỉ là tăng reach. Cần nối nội dung hữu ích với một hành động rõ ràng: nhận guide bằng keyword, "
    "vào DM, xem tour phù hợp, gửi form booking và cuối cùng là đặt dịch vụ."
)

h(doc, "Snapshot Ba Kênh")
make_table(
    doc,
    ["Chỉ số", "Facebook", "Instagram", "Website GA4"],
    [
        ["Quy mô hiện tại", "706 followers", "255 followers", "Dữ liệu 7 ngày"],
        ["Nội dung trả về", "45 bài", "69 media / 68 media trả về", "Trang chủ, tour, gallery, booking"],
        ["Tín hiệu chính", "215 shares", "296.279 views; 124.564 reach", "Khách quốc tế đi sâu vào tour và booking"],
        ["Vai trò đề xuất", "Phân phối và chia sẻ", "Khám phá và tương tác", "Xác nhận ý định và chuyển đổi"],
        ["Giới hạn", "Chưa có comment/reaction", "Metric theo định nghĩa Meta", "Có lẫn test nội bộ và trang quản trị"],
    ],
    [1.6, 1.55, 1.65, 2.2],
)

doc.add_page_break()
h(doc, "1. Hiệu Suất Instagram")
doc.add_paragraph("Cửa sổ Insights truy vấn: 20 tháng 08 đến 18 tháng 09 năm 2026.")
make_table(
    doc,
    ["Chỉ số", "Giá trị hiện tại", "Diễn giải"],
    [
        ["Lượt xem", "296.279", "Khả năng khám phá và lặp lại lượt xem cao"],
        ["Tài khoản tiếp cận", "124.564", "Phễu đầu vào lớn"],
        ["Tài khoản tương tác", "6.888", "Có nền tảng tương tác thực chất"],
        ["Tổng tương tác nội dung", "14.813", "Bao gồm like, save, comment, share hoặc reply"],
        ["Lượt truy cập profile", "2.117", "Cơ hội nối nội dung với offer"],
        ["Lượt bấm link website", "11", "Chuyển đổi website còn thấp"],
        ["Followers", "255", "Tăng nhưng còn nhỏ so với lượng tiếp cận"],
        ["Media", "69", "Kho nội dung đang hoạt động"],
    ],
    [1.8, 1.25, 4.15],
)

h(doc, "Nội Dung Instagram Đang Thắng", 2)
make_table(
    doc,
    ["Nhóm nội dung", "Tín hiệu", "Vì sao hiệu quả"],
    [
        ["Cách sang đường ở Việt Nam", "325 likes, 15 comments", "Giải quyết nỗi lo đầu chuyến đi bằng một mẹo địa phương dễ nhớ"],
        ["Hà Nội ngập và các tình huống bất ngờ", "274 likes, 11 comments", "Kết hợp hài hước với trải nghiệm du lịch thật"],
        ["Địa điểm ẩn ở Hà Nội", "250 likes, 6 comments", "Tạo tò mò và động lực chia sẻ"],
        ["Mẹo tiết kiệm thời gian", "227 likes, 6 comments", "Utility ngắn, dễ lưu"],
        ["Tiền mặt và an toàn đường phố", "159 likes, 19 comments", "Nỗi lo có ý định cao, kèm CTA keyword"],
        ["Danh sách quán cà phê Train Street", "65 likes, 5 comments", "Dạng list hỗ trợ lên lịch và lưu lại"],
    ],
    [2.05, 1.55, 3.6],
)

h(doc, "2. Hiệu Suất Facebook Fanpage")
doc.add_paragraph(
    "Dữ liệu live trả về 706 followers, 45 bài trong khoảng 09–17 tháng 09 năm 2026 và 215 lượt chia sẻ trong tập bài "
    "đọc được. Có 14/45 bài có ít nhất một share; median là 0, cho thấy phân phối tập trung vào một nhóm winner nhỏ."
)
make_table(
    doc,
    ["Hạng", "Mẫu nội dung", "Shares", "Tín hiệu hành động"],
    [
        ["1", "Cách sang đường ở Việt Nam", "66", "Winner utility rõ nhất"],
        ["2", "Vietnam Cheat Sheet Part 2", "28", "Checklist dễ gửi cho bạn đồng hành"],
        ["3", "Địa điểm ẩn ở Hà Nội kèm CTA DM", "28", "Tò mò có thể mở luồng DM"],
        ["4", "10 địa điểm ăn uống ở Hà Nội", "24", "Giá trị lên kế hoạch"],
        ["5", "Mẹo tiết kiệm thời gian", "19", "Ngắn và dễ chuyển tiếp"],
        ["6", "5 điều không nên làm ở Việt Nam", "14", "Hài hước và cảnh báo vẫn ổn định"],
        ["7", "Lỗi khách du lịch thường mắc ở Hà Nội", "13", "Có thể nhân bản theo series"],
    ],
    [0.55, 3.0, 0.8, 2.55],
)

h(doc, "3. Dữ Liệu Website Từ GA4")
doc.add_paragraph(
    "Phần GA4 dưới đây lấy từ kết quả truy vấn Google Analytics do người dùng cung cấp, tập trung vào 7 ngày gần nhất. "
    "Dữ liệu địa lý có thể bao gồm phiên test nội bộ và truy cập trang quản trị, vì vậy cần lọc trước khi dùng làm KPI thương mại."
)
make_table(
    doc,
    ["Khu vực", "Views", "Users", "Thời gian TB", "Tín hiệu chính"],
    [
        ["Hà Nội, Việt Nam", "179", "9", "12 phút", "Trang chủ, tour, booking và trang quản trị nội bộ"],
        ["Kota Kinabalu, Malaysia", "19", "1", "31 phút 42 giây", "Xem /tours 8 lần, gallery và tour Cinematic Walk"],
        ["Hải Phòng, Việt Nam", "9", "1", "25 phút 57 giây", "Tour Chùa và Hồ Tây, gallery, tin tức"],
        ["Pune, Ấn Độ", "6", "1", "49 giây", "Vào /booking 3 lần"],
        ["Essa, Canada", "4", "1", "39 giây", "Vào /booking 2 lần và /how-it-works"],
        ["Madurai, Ấn Độ", "2", "1", "2 phút 55 giây", "Vào /booking và trang chủ"],
        ["Düsseldorf, Đức", "1", "1", "Chưa đủ mẫu", "Trang chủ"],
        ["Dubai, UAE", "1", "1", "Chưa đủ mẫu", "Trang chủ"],
        ["Forest City, Hoa Kỳ", "1", "1", "26 giây", "Trang chủ"],
    ],
    [1.7, 0.65, 0.65, 1.1, 2.8],
)

h(doc, "Tín Hiệu Ý Định Mua", 2)
bullet(doc, "Kota Kinabalu là tín hiệu sâu nhất: 19 lượt xem, hơn 31 phút và đi qua tour, gallery cùng tour Cinematic Hanoi Visual Storytelling Walk.")
bullet(doc, "Pune và Madurai đều chạm trang booking; riêng Pune xem booking 3 lần. Đây là nhóm cần được kiểm tra lại bằng nguồn, UTM và form.")
bullet(doc, "Khách tại Essa, Canada xem /how-it-works và booking, cho thấy nhu cầu tìm hiểu quy trình trước khi liên hệ.")
bullet(doc, "Hải Phòng có thời gian cao nhưng cần tách khách nội địa, team và phiên test trước khi coi là lead.")

h(doc, "4. So Sánh Với Snapshot Trước")
make_table(
    doc,
    ["Chỉ số", "Snapshot trước", "Live hiện tại", "Thay đổi quan sát"],
    [
        ["Facebook followers", "69", "706", "+637"],
        ["Facebook bài đọc được", "29", "45", "+16"],
        ["Facebook top shares", "13", "66", "+53"],
        ["Instagram followers", "191", "255", "+64"],
        ["Instagram media", "59", "69", "+10"],
        ["Instagram views", "219.025", "296.279", "+77.254"],
        ["Instagram reach", "133.224", "124.564", "-8.660"],
        ["Instagram accounts engaged", "5.367", "6.888", "+1.521"],
        ["Instagram interactions", "10.737", "14.813", "+4.076"],
        ["Instagram profile visits", "1.416", "2.117", "+701"],
        ["Instagram website clicks", "7", "11", "+4"],
    ],
    [2.2, 1.45, 1.45, 1.8],
)
doc.add_paragraph(
    "Đây là so sánh định hướng giữa các snapshot được thu thập ở thời điểm và cửa sổ dữ liệu khác nhau, không phải báo cáo "
    "tăng trưởng cùng kỳ tuyệt đối. Instagram tăng views, interactions và profile visits nhưng reach thấp hơn snapshot trước; "
    "Facebook hiện có follower và share cao hơn rõ rệt."
)

h(doc, "5. Kết Luận Liên Kênh Và Phễu Chuyển Đổi")
bullet(doc, "Instagram nên giữ vai trò discovery: Reels utility, nội dung có khả năng save/share và CTA keyword.")
bullet(doc, "Facebook nên tái sử dụng các winner có share cao để mở rộng phân phối và kéo người xem về guide hoặc DM.")
bullet(doc, "Website đã có tín hiệu high-intent từ Malaysia, Ấn Độ và Canada; cần tối ưu landing page, booking và đo UTM.")
bullet(doc, "Nội dung có vấn đề cụ thể của khách du lịch hiệu quả hơn nội dung chỉ đẹp hoặc truyền cảm hứng chung.")
bullet(doc, "Cầu nối thương mại nên đi theo chuỗi: nội dung hữu ích → keyword → DM → guide → chọn nhu cầu → báo giá → booking.")

h(doc, "Kế Hoạch Ưu Tiên")
make_table(
    doc,
    ["Ưu tiên", "Thực hiện", "Đo lường"],
    [
        ["1. Nhân bản winner", "Làm series về sang đường, tiền mặt, lỗi du lịch, food route và weather/packing.", "Shares, saves, comments"],
        ["2. Dùng keyword thống nhất", "HANOI cho guide; PHOTO cho nhu cầu chụp ảnh; trả lời qua DM.", "Keyword comments, DMs"],
        ["3. Sửa website theo ý định", "Đưa tour, gallery, how-it-works và booking thành một luồng rõ ràng.", "Booking views, form starts"],
        ["4. Lọc GA4", "Loại trừ Tag Assistant, trang admin và phiên nội bộ; gắn UTM cho Meta.", "Qualified sessions, source"],
        ["5. Đo doanh thu", "Ghi nguồn lead, quote, booking và giá trị đơn.", "DM-to-booking rate, booking value"],
    ],
    [1.35, 3.75, 1.8],
)

h(doc, "Giới Hạn Dữ Liệu")
bullet(doc, "Facebook chưa trả comments/reactions vì app chưa có pages_read_user_content hoặc Page Public Content Access.")
bullet(doc, "Instagram Insights là số liệu cấp tài khoản theo định nghĩa Meta và có thể là số ước tính.")
bullet(doc, "GA4 bạn cung cấp có lẫn phiên test nội bộ và trang quản trị; không dùng toàn bộ 179 views Hà Nội làm khách thật.")
bullet(doc, "Chưa có dữ liệu booking hoặc doanh thu đủ sạch để kết luận nội dung nào tạo doanh thu cao nhất.")
bullet(doc, "Các mốc trước và hiện tại khác cửa sổ thời gian; phần tăng/giảm chỉ dùng để định hướng, không dùng làm causal proof.")

h(doc, "Kết Luận")
doc.add_paragraph(
    "TravelFrame nên tiếp tục xây tăng trưởng bằng nội dung thực dụng về Hà Nội, đặc biệt là các tình huống khiến khách quốc tế "
    "lo lắng hoặc cần quyết định nhanh. Instagram tạo khám phá và tương tác; Facebook khuếch đại các bài dễ chia sẻ; GA4 cho biết "
    "những khu vực đã có dấu hiệu cân nhắc tour và booking. Mốc đo tiếp theo nên là một chuỗi hoàn chỉnh từ post đến keyword, "
    "DM, lead đủ điều kiện và booking."
)

doc.save(OUT)
print(OUT)
