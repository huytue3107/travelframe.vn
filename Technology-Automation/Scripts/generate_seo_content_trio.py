#!/usr/bin/env python3
"""Build the September 2026 TravelFrame SEO article DOCX deliverables."""

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/macbook/Codex GPT/travelframe.vn")
OUTPUT = ROOT / "Creative-Production" / "Deliverables" / "SEO-Content-Trio-2026-09-25"
ARTICLES = [
    (
        ROOT / "Marketing-Growth/Content/Drafts/Blog-Train-Street-Hanoi-Safety-Guide-2026.md",
        "01-Hanoi-Train-Street-Safety-Guide-2026.docx",
    ),
    (
        ROOT / "Marketing-Growth/Content/Drafts/Blog-Solo-Female-Traveler-Hanoi-Photoshoot.md",
        "02-Solo-Female-Travel-Hanoi-Photo-Guide.docx",
    ),
    (
        ROOT / "Marketing-Growth/Content/Drafts/Blog-Slow-Hanoi-48-Hour-Itinerary.md",
        "03-Slow-Hanoi-48-Hour-Itinerary.docx",
    ),
]

BLACK = RGBColor(0, 0, 0)
BODY = RGBColor(40, 44, 46)
MUTED = RGBColor(94, 104, 108)
TEAL = RGBColor(0, 138, 163)
PALE_TEAL = "EAF5F7"
LIGHT_GRAY = "D9D9D9"


def set_cell_fill(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    props.append(shade)


def set_cell_margins(cell, top=110, start=140, bottom=110, end=140) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell) -> None:
    props = cell._tc.get_or_add_tcPr()
    borders = props.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        props.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), LIGHT_GRAY)


def parse_frontmatter(lines: list[str]) -> tuple[dict[str, str], list[str]]:
    if not lines or lines[0].strip() != "---":
        return {}, lines
    meta: dict[str, str] = {}
    index = 1
    current_list = None
    while index < len(lines):
        line = lines[index]
        if line.strip() == "---":
            return meta, lines[index + 1 :]
        if re.match(r"^\s+-\s+", line) and current_list:
            value = re.sub(r"^\s+-\s+", "", line).strip().strip('"')
            meta[current_list] = f"{meta.get(current_list, '')}, {value}".strip(", ")
        else:
            match = re.match(r"^([a-zA-Z0-9_]+):\s*(.*)$", line)
            if match:
                key, value = match.groups()
                current_list = key
                meta[key] = value.strip().strip('"')
        index += 1
    return meta, lines


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = BODY
    normal.paragraph_format.line_spacing = 1.13
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Heading 3" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(14 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("TravelFrame  |  Draft for review  |  September 25, 2026")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("*", "").replace("`", "")
    return text.strip()


def add_rich_text(paragraph, text: str) -> None:
    link_pattern = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
    bold_pattern = re.compile(r"\*\*([^*]+)\*\*")
    cursor = 0
    tokens = []
    for match in link_pattern.finditer(text):
        if match.start() > cursor:
            tokens.append(("plain", text[cursor : match.start()]))
        tokens.append(("link", f"{match.group(1)} ({match.group(2)})"))
        cursor = match.end()
    if cursor < len(text):
        tokens.append(("plain", text[cursor:]))
    if not tokens:
        tokens = [("plain", text)]

    for kind, chunk in tokens:
        inner_cursor = 0
        for match in bold_pattern.finditer(chunk):
            if match.start() > inner_cursor:
                paragraph.add_run(chunk[inner_cursor : match.start()].replace("*", "").replace("`", ""))
            bold_run = paragraph.add_run(match.group(1))
            bold_run.bold = True
            inner_cursor = match.end()
        if inner_cursor < len(chunk):
            run = paragraph.add_run(chunk[inner_cursor:].replace("*", "").replace("`", ""))
            if kind == "link":
                run.font.color.rgb = TEAL


def add_metadata(doc: Document, meta: dict[str, str]) -> None:
    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    entries = [
        ("SEO slug", meta.get("slug", "")),
        ("Primary keyword", meta.get("primary_keyword", "")),
        ("Status", meta.get("status", "Draft")),
        ("Audience", meta.get("target_persona", "")),
        ("Meta title", meta.get("meta_title", "")),
        ("Meta description", meta.get("meta_description", "")),
    ]
    for idx, (label, value) in enumerate(entries):
        row, col = divmod(idx, 2)
        cell = table.rows[row].cells[col]
        set_cell_fill(cell, PALE_TEAL if row % 2 == 0 else "FFFFFF")
        set_cell_margins(cell)
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        label_run = p.add_run(f"{label}\n")
        label_run.bold = True
        label_run.font.size = Pt(8.5)
        label_run.font.color.rgb = MUTED
        value_run = p.add_run(value)
        value_run.font.size = Pt(9.2)
        value_run.font.color.rgb = BODY
    doc.add_paragraph()


def resolve_image(md_file: Path, reference: str) -> Path:
    return (md_file.parent / reference).resolve()


def add_image(doc: Document, path: Path, alt: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.55))
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))
    paragraph._travel_frame_alt = alt


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.rows[row_index].cells[col_index]
            set_cell_fill(cell, "314F59" if row_index == 0 else ("F3F8F9" if row_index % 2 == 0 else "FFFFFF"))
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(clean_inline(value))
            run.font.name = "Aptos"
            run.font.size = Pt(9)
            run.font.bold = row_index == 0
            run.font.color.rgb = RGBColor(255, 255, 255) if row_index == 0 else BODY
    doc.add_paragraph()


def build_article(doc: Document, md_file: Path, include_title=True) -> str:
    raw_lines = md_file.read_text(encoding="utf-8").splitlines()
    meta, lines = parse_frontmatter(raw_lines)
    title = meta.get("title", md_file.stem)
    if include_title:
        title_paragraph = doc.add_paragraph(style="Title")
        title_paragraph.add_run(title)
        intro = doc.add_paragraph("SEO article draft with curated Pexels photography and clearly disclosed AI editorial illustration.")
        intro.paragraph_format.space_after = Pt(12)
        intro.runs[0].font.color.rgb = MUTED
        intro.runs[0].font.size = Pt(10.5)
        add_metadata(doc, meta)

    in_code = False
    table_rows: list[list[str]] = []
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            index += 1
            continue
        if not stripped or stripped == "---":
            if table_rows:
                add_table(doc, table_rows)
                table_rows = []
            index += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            values = [part.strip() for part in stripped[1:-1].split("|")]
            if not all(re.fullmatch(r":?-+:?", value) for value in values):
                table_rows.append(values)
            index += 1
            continue
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            add_image(doc, resolve_image(md_file, image_match.group(2)), image_match.group(1))
            index += 1
            continue

        if stripped.startswith("# "):
            index += 1
            continue
        if stripped.startswith("## "):
            doc.add_paragraph(clean_inline(stripped[3:]), style="Heading 1")
        elif stripped.startswith("### "):
            doc.add_paragraph(clean_inline(stripped[4:]), style="Heading 2")
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_rich_text(p, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_rich_text(p, stripped[2:])
        elif stripped.startswith("*") and stripped.endswith("*"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(10)
            run = p.add_run(clean_inline(stripped))
            run.italic = True
            run.font.size = Pt(8.5)
            run.font.color.rgb = MUTED
        else:
            p = doc.add_paragraph()
            add_rich_text(p, stripped)
        index += 1
    if table_rows:
        add_table(doc, table_rows)
    return title


def save_individuals() -> list[tuple[str, Path]]:
    outputs = []
    for md_file, filename in ARTICLES:
        doc = Document()
        configure_document(doc)
        title = build_article(doc, md_file)
        path = OUTPUT / filename
        doc.save(path)
        outputs.append((title, path))
    return outputs


def save_master() -> Path:
    doc = Document()
    configure_document(doc)
    title = doc.add_paragraph(style="Title")
    title.add_run("TravelFrame SEO Content Trio")
    subtitle = doc.add_paragraph("Three English SEO article drafts with visual assets and source notes")
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(18)

    summary = doc.add_table(rows=4, cols=3)
    summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ("Article", "Primary intent", "Visuals"),
        ("Hanoi Train Street 2026", "Safety and access", "1 Codex illustration + 2 Pexels photos"),
        ("Solo Female Travel in Hanoi", "Photo guidance and consideration", "1 Codex illustration + 2 Pexels photos"),
        ("Slow Hanoi 48-Hour Itinerary", "Two-day travel planning", "1 Codex illustration + 3 Pexels photos"),
    ]
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = summary.rows[r_idx].cells[c_idx]
            set_cell_fill(cell, "314F59" if r_idx == 0 else ("F3F8F9" if r_idx % 2 == 0 else "FFFFFF"))
            set_cell_margins(cell)
            set_cell_borders(cell)
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(9)
            run.font.bold = r_idx == 0
            run.font.color.rgb = RGBColor(255, 255, 255) if r_idx == 0 else BODY
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(14)
    note.add_run("Commercial claims: ").bold = True
    note.add_run("Prices, delivery promises, image counts and testimonials were omitted because they are not approved commercial truth.")

    for index, (md_file, _) in enumerate(ARTICLES):
        doc.add_page_break()
        build_article(doc, md_file, include_title=True)
        if index < len(ARTICLES) - 1:
            doc.add_section(WD_SECTION.CONTINUOUS)

    path = OUTPUT / "TravelFrame-SEO-Content-Trio-2026-09-25.docx"
    doc.save(path)
    return path


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    for _, old_path in ARTICLES:
        _ = old_path
    outputs = save_individuals()
    master = save_master()
    print(master)
    for _, path in outputs:
        print(path)


if __name__ == "__main__":
    main()
