from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]

OUT = ROOT / "Data-Analytics" / "Reports" / "2026-09-18-TravelFrame-Meta-Combined-Report.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

TEAL = "008AA3"
DARK = "222222"
LIGHT_BLUE = "EAF5F7"
LIGHT_GRAY = "F4F6F7"
BORDER = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, color=DARK, size=9.5, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        set_cell_shading(cell, TEAL)
        set_cell_text(cell, header, bold=True, color="FFFFFF", size=9)
        if widths:
            cell.width = Inches(widths[i])
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            if widths:
                cells[i].width = Inches(widths[i])
            set_cell_shading(cells[i], "FFFFFF" if ridx % 2 == 0 else LIGHT_GRAY)
            set_cell_text(cells[i], value, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor.from_string(DARK)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for style_name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12.5), ("Heading 3", 11)):
    style = styles[style_name]
    style.font.name = "Aptos Display" if style_name == "Title" else "Aptos"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string("000000")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("TravelFrame | Meta Channel Report | 18 September 2026")
footer_run.font.name = "Aptos"
footer_run.font.size = Pt(8)
footer_run.font.color.rgb = RGBColor.from_string("666666")

title = doc.add_paragraph(style="Title")
title.add_run("TravelFrame Meta Channel Report")
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run("Facebook Page and Instagram performance comparison")
run.font.name = "Aptos"
run.font.size = Pt(14)
run.font.color.rgb = RGBColor.from_string("000000")

meta = table(
    doc,
    ["Field", "Value"],
    [
        ["Report date", "18 September 2026"],
        ["Facebook Page", "travelframe.vn | Page ID 1409099405611405"],
        ["Instagram", "@travelframe.vn | Business Account ID 17841438272497512"],
        ["Data source", "Meta Graph API via read-only MCP and Page-token smoke check"],
        ["Status", "Live read / Draft for review"],
    ],
    widths=[1.45, 5.45],
)

heading(doc, "Executive Summary", 1)
doc.add_paragraph(
    "TravelFrame is currently strongest as a practical Hanoi and Vietnam travel utility brand. "
    "Instagram is the larger discovery engine, while Facebook is showing a sharp improvement in "
    "Page scale and shareable distribution. The clearest cross-channel winner is content that removes "
    "travel anxiety or gives visitors a short, useful action they can send to a travel companion."
)
doc.add_paragraph(
    "The immediate growth opportunity is to keep the utility content that earns reach and shares, then "
    "add a consistent bridge to TravelFrame's local guide and photography services. The data supports "
    "this direction, but it does not yet prove which channel or post caused follower growth or bookings."
)

heading(doc, "Live Channel Snapshot", 1)
table(
    doc,
    ["Metric", "Facebook Page", "Instagram", "Reading"],
    [
        ["Followers", "706", "255", "Facebook is larger in current follower count"],
        ["Published media/posts returned", "45 posts", "69 profile media / 68 returned", "Instagram has the deeper content library"],
        ["Latest live content date", "17 Sep 2026", "17 Sep 2026", "Both channels are publishing recently"],
        ["Primary signal available", "Shares", "Views, reach, interactions, profile visits", "Metrics are not identical across platforms"],
        ["Main limitation", "No comments/reactions", "Insights are estimated and API-defined", "Do not treat channel totals as directly interchangeable"],
    ],
    widths=[1.65, 1.7, 1.85, 2.0],
)

doc.add_page_break()
heading(doc, "Current Instagram Performance", 1)
doc.add_paragraph("Insights query window: 20 August to 18 September 2026. Meta returned these account-level totals:")
table(
    doc,
    ["Metric", "Current value", "Interpretation"],
    [
        ["Views", "296,279", "Strong discovery and repeat exposure"],
        ["Accounts reached", "124,564", "Large top-of-funnel audience"],
        ["Accounts engaged", "6,888", "Useful interaction base relative to reach"],
        ["Content interactions", "14,813", "Likes, saves, comments, shares or replies"],
        ["Profile visits", "2,117", "A meaningful bridge opportunity"],
        ["Website link taps", "11", "Website conversion remains very small"],
        ["Followers", "255", "Follower base has grown but remains modest versus reach"],
        ["Media count", "69", "Active publishing library"],
    ],
    widths=[1.7, 1.25, 4.25],
)

heading(doc, "Instagram Content Winners", 2)
table(
    doc,
    ["Content pattern", "Observed signal", "Why it works"],
    [
        ["Crossing the road", "325 likes, 15 comments", "Solves a first-trip anxiety with a memorable local insight"],
        ["Hanoi flooding / unexpected city moments", "274 likes, 11 comments", "Humor plus a recognizable travel experience"],
        ["Hidden Hanoi location", "250 likes, 6 comments", "Curiosity and location value create sharing intent"],
        ["Save time travel tip", "227 likes, 6 comments", "Short, practical utility"],
        ["Cash and street safety", "159 likes, 19 comments", "High-intent concern with clear advice and keyword CTA"],
        ["Train Street cafés", "65 likes, 5 comments", "List format supports saving and route planning"],
    ],
    widths=[2.05, 1.55, 3.6],
)

doc.add_page_break()
heading(doc, "Current Facebook Page Performance", 1)
doc.add_paragraph(
    "The live Page read returned 706 followers, 45 posts dated 9–17 September 2026, and 215 total "
    "shares across the returned set. Fourteen of the 45 posts had at least one share; the median was zero, "
    "so distribution is concentrated in a small group of winners."
)
table(
    doc,
    ["Rank", "Post pattern", "Shares", "Actionable signal"],
    [
        ["1", "Crossing the road in Vietnam", "66", "Strongest current utility and share trigger"],
        ["2", "Vietnam Cheat Sheet Part 2", "28", "Checklist format travels well between friends"],
        ["3", "Hidden Hanoi location with DM CTA", "28", "Curiosity can open a DM path"],
        ["4", "10 Hanoi food spots", "24", "List format creates planning value"],
        ["5", "Save time travel tip", "19", "Short utility is easy to forward"],
        ["6", "5 Things Not To Do In Vietnam", "14", "Humor and warnings remain reliable"],
        ["7", "Tourist mistakes in Hanoi", "13", "Pain-point framing is repeatable"],
    ],
    widths=[0.55, 3.0, 0.8, 2.55],
)

heading(doc, "Comparison With Previous Snapshots", 1)
doc.add_paragraph(
    "The previous recorded snapshot reported 69 Facebook followers, 29 readable posts from 9–14 September, "
    "and a top observed post at 13 shares. The previous Instagram API snapshot reported 191 followers, 59 media, "
    "219,025 views, 133,224 reach, 5,367 accounts engaged, 10,737 interactions, 1,416 profile visits and 7 website clicks."
)
table(
    doc,
    ["Metric", "Previous snapshot", "Current live", "Observed change"],
    [
        ["Facebook followers", "69", "706", "+637"],
        ["Facebook readable posts", "29", "45", "+16"],
        ["Facebook top shares", "13", "66", "+53"],
        ["Instagram followers", "191", "255", "+64"],
        ["Instagram media", "59", "69", "+10"],
        ["Instagram views", "219,025", "296,279", "+77,254"],
        ["Instagram reach", "133,224", "124,564", "-8,660"],
        ["Instagram accounts engaged", "5,367", "6,888", "+1,521"],
        ["Instagram interactions", "10,737", "14,813", "+4,076"],
        ["Instagram profile visits", "1,416", "2,117", "+701"],
        ["Instagram website clicks", "7", "11", "+4"],
    ],
    widths=[2.2, 1.45, 1.45, 1.8],
)
doc.add_paragraph(
    "Interpretation: Instagram shows higher views, interactions and profile visits in the latest query, while "
    "reach is lower than the previous snapshot. Facebook shows a much larger current follower count and stronger "
    "share performance than the earlier Page snapshot. These are directional comparisons, not strict like-for-like "
    "period growth rates, because the extraction dates and API windows differ."
)

heading(doc, "Cross Channel Findings", 1)
bullet(doc, "The same practical topic can travel across both platforms: crossing roads, cash safety, tourist mistakes and Hanoi lists.")
bullet(doc, "Instagram is better positioned for discovery and profile action; Facebook is currently showing stronger follower scale and share concentration.")
bullet(doc, "Content with a clear travel problem performs better than aesthetic-only content.")
bullet(doc, "The strongest CTA pattern is a keyword that unlocks a useful guide, not a hard sell in the first sentence.")
bullet(doc, "The current funnel gap remains the bridge from useful travel content to a paid photo-tour or local experience.")

heading(doc, "Recommended Content and Conversion Plan", 1)
table(
    doc,
    ["Priority", "Execution", "Measurement"],
    [
        ["1. Multiply winners", "Create a 5–7 post series around road crossing, cash safety, food routes and tourist mistakes.", "Shares, saves, comments"],
        ["2. Keep a utility library", "Publish cheat sheets, weather/packing notes and short neighborhood routes.", "Reach, profile visits"],
        ["3. Use one keyword bridge", "Use HANOI for guides and PHOTO for photography intent; route replies into DM.", "Keyword comments, DMs"],
        ["4. Show the paid offer softly", "After utility, show how a local host and photographer make the route easier and more memorable.", "Qualified leads, quotes"],
        ["5. Track the whole funnel", "Record source, DM, qualified lead, quote and paid booking.", "DM-to-booking rate, booking value"],
    ],
    widths=[1.35, 3.75, 1.8],
)

heading(doc, "Data Limitations", 1)
bullet(doc, "Facebook comments and reactions were not returned because pages_read_user_content or Page Public Content Access is not enabled.")
bullet(doc, "Facebook totals in this report are based on the returned Page post set, not a full Page Insights export.")
bullet(doc, "Instagram Insights are Meta API account-level totals and may be estimated under Meta's metric definitions.")
bullet(doc, "The previous and current snapshots were collected at different times and should not be treated as identical reporting periods.")
bullet(doc, "No booking or revenue data was available in the channel read, so content performance is not the same as commercial performance.")

heading(doc, "Conclusion", 1)
doc.add_paragraph(
    "TravelFrame should continue building its audience through practical Hanoi travel help, with crossing-the-road "
    "content as the clearest current test case. Instagram should carry discovery and keyword-led profile action; "
    "Facebook should reuse the most shareable utility posts and support distribution. The next measurement milestone "
    "is not more reach alone: it is a clean chain from post to keyword response, DM, qualified lead and booking."
)

doc.save(OUT)
print(OUT)
