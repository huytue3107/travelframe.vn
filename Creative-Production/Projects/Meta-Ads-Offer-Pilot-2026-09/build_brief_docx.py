from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
PROJECT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Creative-Production" / "Deliverables" / "Docx"

BRAND_RED = "AE2A1D"
DEEP_RED = "42100E"
YELLOW = "FFE050"
BLACK = "101010"
WHITE = "FFFFFF"
WARM_WHITE = "F7F2ED"
LIGHT_GRAY = "EDEDED"
MID_GRAY = "D9D9D9"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), MID_GRAY)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, enabled=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    keep = p_pr.find(qn("w:keepNext"))
    if enabled and keep is None:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    elif not enabled and keep is not None:
        p_pr.remove(keep)


def remove_paragraph_borders(paragraph_or_style) -> None:
    element = paragraph_or_style._p if hasattr(paragraph_or_style, "_p") else paragraph_or_style.element
    p_pr = element.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_inline_runs(paragraph, text: str, base_size: float = 10.8) -> None:
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            run.font.size = Pt(base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(base_size)
        else:
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Arial"
            run.font.size = Pt(base_size - 0.3)
            run.font.color.rgb = RGBColor.from_string(DEEP_RED)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        run.font.size = Pt(base_size)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.14

    title = styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    remove_paragraph_borders(title)

    for name, size, before, after in (
        ("Heading 1", 17, 16, 7),
        ("Heading 2", 14, 13, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Arial"
        style.font.size = Pt(10.6)
        style.paragraph_format.space_after = Pt(3)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    header = section.header.paragraphs[0]
    header.text = "TRAVELFRAME   CREATIVE PRODUCTION"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(BRAND_RED)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(BLACK)


def add_cover(doc: Document, title: str, subtitle: str, recipient: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("TRAVELFRAME")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BRAND_RED)

    title_p = doc.add_paragraph(style="Title")
    title_p.add_run(title)
    remove_paragraph_borders(title_p)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(20)
    run = subtitle_p.add_run(subtitle)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(BLACK)

    metadata = doc.add_table(rows=4, cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.LEFT
    metadata.autofit = False
    metadata.columns[0].width = Inches(1.55)
    metadata.columns[1].width = Inches(5.35)
    values = [
        ("Campaign", "TF LEADS HANOI OFFER TEST 202609"),
        ("Recipient", recipient),
        ("Status", "Review  Chưa publish hoặc chạy ads"),
        ("Reference date", "15 September 2026"),
    ]
    for idx, (label, value) in enumerate(values):
        left, right = metadata.rows[idx].cells
        set_cell_fill(left, BRAND_RED)
        set_cell_fill(right, WARM_WHITE if idx % 2 == 0 else WHITE)
        for cell in (left, right):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        left.text = ""
        label_run = left.paragraphs[0].add_run(label)
        label_run.bold = True
        label_run.font.name = "Arial"
        label_run.font.size = Pt(9.5)
        label_run.font.color.rgb = RGBColor.from_string(WHITE)
        right.text = ""
        value_run = right.paragraphs[0].add_run(value)
        value_run.font.name = "Arial"
        value_run.font.size = Pt(9.5)
        value_run.font.color.rgb = RGBColor.from_string(BLACK)
    set_table_borders(metadata)

    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.paragraph_format.space_before = Pt(8)
    intro.paragraph_format.space_after = Pt(10)
    add_inline_runs(
        intro,
        "Tài liệu này chuyển campaign pilot 300 USD thành yêu cầu sản xuất có thể giao "
        "việc trực tiếp. Mỗi creative phải làm rõ một offer trong ba giây đầu, dùng "
        "đúng nhận diện Facebook và Instagram, và dẫn người xem tới một keyword DM.",
        11.2,
    )

    palette = doc.add_table(rows=1, cols=5)
    palette.alignment = WD_TABLE_ALIGNMENT.LEFT
    palette.autofit = False
    colors = [
        ("Brand red", BRAND_RED, WHITE),
        ("Deep red", DEEP_RED, WHITE),
        ("Highlight", YELLOW, BLACK),
        ("Black", BLACK, WHITE),
        ("White", WHITE, BLACK),
    ]
    for idx, (label, fill, text_color) in enumerate(colors):
        cell = palette.rows[0].cells[idx]
        set_cell_fill(cell, fill)
        set_cell_margins(cell, top=150, bottom=150, start=80, end=80)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"{label}\n#{fill}")
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(text_color)
    set_table_borders(palette)

    doc.add_page_break()


def add_markdown_table(doc: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c) for c in rows[1]):
        rows.pop(1)
    if not rows:
        return

    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)

    for row_idx, values in enumerate(rows):
        row = table.rows[row_idx]
        set_cant_split(row)
        if row_idx == 0:
            set_repeat_table_header(row)
        for col_idx in range(width):
            cell = row.cells[col_idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = values[col_idx] if col_idx < len(values) else ""
            cell.text = ""
            paragraph = cell.paragraphs[0]
            if row_idx == 0:
                set_cell_fill(cell, DEEP_RED)
                run = paragraph.add_run(value)
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(WHITE)
            else:
                set_cell_fill(cell, WARM_WHITE if row_idx % 2 == 0 else WHITE)
                add_inline_runs(paragraph, value, 9.1)
            for run in paragraph.runs:
                run.font.name = "Arial"
                if run.font.size is None:
                    run.font.size = Pt(9.1)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9.2)
        run.font.color.rgb = RGBColor.from_string(DEEP_RED)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def should_page_break(heading: str, source_kind: str) -> bool:
    heading = re.sub(r"^\d+\.\s*", "", heading)
    if source_kind == "video":
        return heading.startswith(("Creative 01", "Creative 02", "Creative 03", "Review Và Feedback"))
    return heading.startswith(
        (
            "Offer 01 - FILM Static",
            "Offer 01 - FILM Carousel",
            "Offer 02 - PHOTO Static",
            "Offer 02 - PHOTO Carousel",
            "Offer 03 - HANOI Static",
            "Offer 03 - HANOI Carousel",
            "Review Process",
        )
    )


def add_markdown_body(doc: Document, source: Path, source_kind: str) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    index = next(
        (idx for idx, line in enumerate(lines) if re.match(r"^##\s+1\.", line.strip())),
        1,
    )
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(raw)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped.startswith("**") and stripped.endswith("**") and ":" in stripped:
            p = doc.add_paragraph()
            add_inline_runs(p, stripped, 10)
            index += 1
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1)) - 1
            heading = heading_match.group(2)
            p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            if should_page_break(heading, source_kind):
                p.paragraph_format.page_break_before = True
            add_inline_runs(p, heading, {1: 17, 2: 14, 3: 11.5}[min(level, 3)])
            set_keep_with_next(p)
            index += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(doc, table_lines)
            continue

        checklist_match = re.match(r"^-\s+\[([ xX])\]\s+(.+)$", stripped)
        if checklist_match:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(3)
            symbol = "☒" if checklist_match.group(1).lower() == "x" else "☐"
            add_inline_runs(p, f"{symbol} {checklist_match.group(2)}", 10.5)
            index += 1
            continue

        bullet_match = re.match(r"^-\s+(.+)$", stripped)
        if bullet_match:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.20)
            p.paragraph_format.space_after = Pt(3)
            add_inline_runs(p, f"•  {bullet_match.group(1)}", 10.6)
            index += 1
            continue

        number_match = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if number_match:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(3)
            add_inline_runs(p, f"{number_match.group(1)}.  {number_match.group(2)}", 10.6)
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            upcoming = lines[index].strip()
            if (
                not upcoming
                or upcoming.startswith(("#", "|", "-", "```"))
                or re.match(r"^\d+\.\s+", upcoming)
            ):
                break
            paragraph_lines.append(upcoming)
            index += 1
        p = doc.add_paragraph()
        add_inline_runs(p, " ".join(paragraph_lines), 10.8)

    if code_lines:
        add_code_block(doc, code_lines)


def build_document(source_name: str, output_name: str, title: str, subtitle: str, recipient: str, source_kind: str) -> Path:
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_cover(doc, title, subtitle, recipient)
    add_markdown_body(doc, PROJECT / source_name, source_kind)

    core = doc.core_properties
    core.title = title
    core.subject = subtitle
    core.author = "TravelFrame"
    core.keywords = "TravelFrame, Meta Ads, creative brief, Hanoi"

    OUTPUT.mkdir(parents=True, exist_ok=True)
    output = OUTPUT / output_name
    doc.save(output)
    return output


def main() -> None:
    outputs = [
        build_document(
            "CREATIVE-PRODUCTION-BRIEF.md",
            "TravelFrame-Video-Creative-Production-Brief-2026-09-15.docx",
            "TravelFrame Video Creative Production Brief",
            "Production requirements for the FILM PHOTO and HANOI Meta Ads creatives",
            "Graphic Designer and Video Editor",
            "video",
        ),
        build_document(
            "STATIC-AND-CAROUSEL-DESIGN-BRIEF.md",
            "TravelFrame-Static-Carousel-Design-Brief-2026-09-15.docx",
            "TravelFrame Static Post and Carousel Design Brief",
            "Design requirements for the FILM PHOTO and HANOI Meta Ads creatives",
            "Graphic Designer and Content Marketing",
            "static",
        ),
    ]
    for output in outputs:
        print(output)


if __name__ == "__main__":
    main()
