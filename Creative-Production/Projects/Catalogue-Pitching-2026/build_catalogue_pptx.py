from __future__ import annotations

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[3]
OUTPUT = ROOT / "Creative-Production" / "Deliverables" / "PPTX" / "TravelFrame-Customer-Pitching-Catalogue-2026.pptx"
ASSETS = ROOT / "Creative-Production" / "Deliverables" / "Catalogue" / "assets"

TEAL = RGBColor(0, 138, 163)
DARK_TEAL = RGBColor(0, 89, 107)
NAVY = RGBColor(15, 23, 42)
SLATE = RGBColor(30, 41, 59)
MUTED = RGBColor(100, 116, 139)
LIGHT_BG = RGBColor(241, 245, 249)
WHITE = RGBColor(255, 255, 255)
BORDER_COLOR = RGBColor(226, 232, 240)


def create_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def add_blank_slide(prs: Presentation):
    blank_layout = prs.slide_layouts[6]
    return prs.slides.add_slide(blank_layout)


def set_slide_background(slide, color: RGBColor):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, badge: str, title: str, subtitle: str = ""):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(1.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_badge = tf.paragraphs[0]
    p_badge.text = badge.upper()
    p_badge.font.name = "Arial"
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = TEAL

    p_title = tf.add_paragraph()
    p_title.text = title
    p_title.font.name = "Arial"
    p_title.font.size = Pt(21)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY
    p_title.space_before = Pt(3)

    if subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(11.5)
        p_sub.font.color.rgb = MUTED
        p_sub.space_before = Pt(2)


# ==========================================
# SLIDE 1: TRANG BÌA (COVER)
# ==========================================
def build_slide_1(prs: Presentation):
    slide = add_blank_slide(prs)
    set_slide_background(slide, NAVY)

    hero_img = ASSETS / "hero-cover.jpg"
    if hero_img.exists():
        slide.shapes.add_picture(str(hero_img), Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))

    overlay = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.1), Inches(10.9), Inches(5.3))
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = NAVY
    overlay.line.color.rgb = TEAL
    overlay.line.width = Pt(1.5)

    tf = overlay.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.7)

    p0 = tf.paragraphs[0]
    p0.text = "TRAVELFRAME HANOI — BẢN GIỚI THIỆU DỊCH VỤ 2026"
    p0.font.name = "Arial"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(56, 189, 248)

    p1 = tf.add_paragraph()
    p1.text = "Khám Phá Hà Nội Cùng Bạn Bản Địa,\nLưu Giữ Kỷ Niệm Qua Từng Khung Hình Điện Ảnh"
    p1.font.name = "Arial"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_before = Pt(12)

    p2 = tf.add_paragraph()
    p2.text = "Tour đi bộ riêng tư kết hợp cùng dịch vụ chụp ảnh nghệ thuật & quay video ngắn 4K Reels/TikTok dành riêng cho du khách."
    p2.font.name = "Arial"
    p2.font.size = Pt(13.5)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.space_before = Pt(12)

    p3 = tf.add_paragraph()
    p3.text = "✨ 100% Nhịp Đi Riêng Tư   •   ⚡ 24h Nhận Ảnh Xem Trước   •   📍 Tuyến Phố Ẩn Giấu   •   ☕ Cà Phê Đặc Sản Trọn Gói"
    p3.font.name = "Arial"
    p3.font.size = Pt(11)
    p3.font.bold = True
    p3.font.color.rgb = TEAL
    p3.space_before = Pt(18)


# ==========================================
# SLIDE 2: SỰ KHÁC BIỆT CỦA TRAVELFRAME
# ==========================================
def build_slide_2(prs: Presentation):
    slide = add_blank_slide(prs)
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Điểm Khác Biệt Cốt Lõi", "Tạm Biệt Những Buổi Chụp Gượng Gạo & Tour Ghép Vội Vã", "TravelFrame giải quyết triệt để sự thỏa hiệp mà du khách thường gặp phải khi đến Hà Nội.")

    c1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.9), Inches(5.6), Inches(5.0))
    c1.fill.solid()
    c1.fill.fore_color.rgb = WHITE
    c1.line.color.rgb = BORDER_COLOR
    tf1 = c1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_right = tf1.margin_top = Inches(0.4)

    p1 = tf1.paragraphs[0]
    p1.text = "❌ Nỗi Thất Vọng Khi Đi Tour Thông Thường"
    p1.font.name = "Arial"
    p1.font.size = Pt(15)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(220, 38, 38)

    items1 = [
        "• Tour Ghép Đông Đúc: Lịch trình cứng nhắc, dừng 15 phút vội vàng, bị ép vào các điểm bán đồ lưu niệm.",
        "• Chụp Ảnh Gượng Gạo: Thợ ảnh studio lạ lẫm, bắt tạo dáng đơ cứng, không hiểu câu chuyện văn hóa Hà Nội.",
        "• Vất Vả Tự Dựng Tripod: Tự xoay xở chân máy giữa ngõ phố đông người, dễ mất an toàn và bỏ lỡ khoảnh khắc.",
        "• Chờ Ảnh Quá Lâu: Hàng tuần sau chuyến đi mới nhận được file ảnh gốc chưa qua chọn lọc và chỉnh màu."
    ]
    for it in items1:
        p = tf1.add_paragraph()
        p.text = it
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = SLATE
        p.space_before = Pt(12)

    c2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.9), Inches(5.7), Inches(5.0))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE
    c2.line.color.rgb = TEAL
    c2.line.width = Pt(2.5)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = tf2.margin_top = Inches(0.4)

    p2 = tf2.paragraphs[0]
    p2.text = "✨ Trải Nghiệm Độc Bản Cùng TravelFrame"
    p2.font.name = "Arial"
    p2.font.size = Pt(15)
    p2.font.bold = True
    p2.font.color.rgb = DARK_TEAL

    items2 = [
        "• Người Bạn Bản Địa Tri Thức: Host tiếng Anh lưu loát, am hiểu văn hóa ngõ phố, vừa kể chuyện vừa đồng hành.",
        "• Khoảnh Khắc Tự Nhiên (Candid): Hướng dẫn tư thế nhẹ nhàng, thoải mái trò chuyện để bắt trọn nụ cười tự nhiên nhất.",
        "• Thiết Bị Full-Frame Chuyên Nghiệp: Máy ảnh Sony/Canon ống kính chân dung xóa phông mịn màng và mic thu âm 4K.",
        "• 24h Nhận Ảnh Sneak Peek: Nhận ngay 3-5 tấm ảnh hero xuất sắc nhất vào ngày hôm sau để đăng mạng xã hội ngay!"
    ]
    for it in items2:
        p = tf2.add_paragraph()
        p.text = it
        p.font.name = "Arial"
        p.font.size = Pt(11.5)
        p.font.color.rgb = SLATE
        p.space_before = Pt(12)


# ==========================================
# SLIDE 3: 4 PHONG CÁCH HÀ NỘI
# ==========================================
def build_slide_3(prs: Presentation):
    slide = add_blank_slide(prs)
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Định Hướng Thẩm Mỹ", "Bốn Phong Cách Trải Nghiệm & Hình Ảnh Đặc Trưng", "Du khách tự do lựa chọn 1 hoặc 2 phong cách yêu thích. TravelFrame sẽ tùy biến lịch trình riêng cho bạn.")

    moods = [
        ("heritage.jpg", "🏛️ Heritage Hanoi", "Mái ngói đền chùa cổ kính, bức tường rêu phong Văn Miếu, chiều sâu kiến trúc ngàn năm.", Inches(0.8)),
        ("everyday.jpg", "☕ Everyday Hanoi", "Chợ truyền thống Đồng Xuân, ngõ cà phê trứng, Train Street an toàn, nhịp sống hè phố chân thật.", Inches(3.8)),
        ("in-style.jpg", "✨ Hanoi In Style", "Khu phố Pháp thanh lịch, Nhà hát Lớn, Metropole, đại lộ rợp bóng cây Phan Đình Phùng.", Inches(6.8)),
        ("after-dark.jpg", "🌙 Hanoi After Dark", "Hoàng hôn trên cầu Long Biên, phố đêm Tạ Hiện, ánh đèn neon phản chiếu trên mặt hồ.", Inches(9.8)),
    ]

    for img_name, title, desc, left in moods:
        img_path = ASSETS / img_name
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.9), Inches(2.7), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_COLOR

        if img_path.exists():
            slide.shapes.add_picture(str(img_path), left, Inches(1.9), width=Inches(2.7), height=Inches(2.4))

        tb = slide.shapes.add_textbox(left + Inches(0.15), Inches(4.45), Inches(2.4), Inches(2.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(13)
        p1.font.bold = True
        p1.font.color.rgb = DARK_TEAL

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.name = "Arial"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = MUTED
        p2.space_before = Pt(6)


# ==========================================
# SLIDE 4: CÁC GÓI DỊCH VỤ TRỌN GÓI
# ==========================================
def build_slide_4(prs: Presentation):
    slide = add_blank_slide(prs)
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Bảng Giá Minh Bạch", "Các Gói Trải Nghiệm Riêng Tư Trọn Gói (All-Inclusive)", "Giá trọn gói cho cả nhóm riêng tư. Đã bao gồm Host/Nhiếp ảnh gia, vé tham quan, nước uống và chỉnh màu ảnh.")

    pkgs = [
        (
            "1. Dạo Phố Nhanh (The Quick Photo Walk)",
            "⏱️ 2 - 2.5 Giờ (Nhóm 1 - 2 Khách)",
            "1.950.000 VNĐ",
            "$75 USD / Nhóm",
            ["• 15+ Ảnh retouch nghệ thuật", "• Toàn bộ ảnh gốc chất lượng cao", "• 1 Ly cà phê đặc sản Hà Nội", "• 1 Khu vực trọng tâm (Phố Cổ/Phố Pháp)", "• Bàn giao ảnh trong 24 - 48h"],
            False,
            Inches(0.8)
        ),
        (
            "2. Hành Trình Câu Chuyện (Signature Story Tour)",
            "⏱️ 4 - 4.5 Giờ (Nhóm 1 - 3 Khách)",
            "4.150.000 VNĐ",
            "$160 USD / Nhóm",
            ["• 35+ Ảnh retouch cao cấp", "• 1 Video Reel dọc 9:16 (45s) sẵn nhạc", "• Thưởng thức ẩm thực & cà phê trứng", "• Kết hợp 2 Moods thẩm mỹ tùy chọn", "• ⚡ 24h Nhận Sneak Peek + 48h Gallery"],
            True,
            Inches(4.8)
        ),
        (
            "3. Khám Phá Toàn Diện (Cultural Odyssey)",
            "⏱️ 8 - 9 Giờ (Nhóm 2 - 6 Khách)",
            "7.500.000 VNĐ",
            "$290 USD / Nhóm",
            ["• 60+ Ảnh hoàn thiện chuẩn tạp chí", "• 1 Phim ngắn cinematic (60-90s)", "• Xe hơi riêng đời mới trọn ngày", "• Làng gốm Bát Tràng hoặc Làng hương", "• Bữa trưa ẩm thực truyền thống"],
            False,
            Inches(8.8)
        ),
    ]

    for title, duration, vnd, usd, features, popular, left in pkgs:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.9), Inches(3.7), Inches(5.0))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(251, 254, 255) if popular else WHITE
        card.line.color.rgb = TEAL if popular else BORDER_COLOR
        card.line.width = Pt(2.5) if popular else Pt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.3)

        if popular:
            p_pop = tf.paragraphs[0]
            p_pop.text = "⭐ LỰA CHỌN PHỔ BIẾN NHẤT"
            p_pop.font.name = "Arial"
            p_pop.font.size = Pt(10)
            p_pop.font.bold = True
            p_pop.font.color.rgb = TEAL
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        p_title.text = title
        p_title.font.name = "Arial"
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

        p_dur = tf.add_paragraph()
        p_dur.text = duration
        p_dur.font.name = "Arial"
        p_dur.font.size = Pt(11)
        p_dur.font.color.rgb = MUTED
        p_dur.space_before = Pt(4)

        p_pr = tf.add_paragraph()
        p_pr.text = f"{vnd}  ({usd})"
        p_pr.font.name = "Arial"
        p_pr.font.size = Pt(15)
        p_pr.font.bold = True
        p_pr.font.color.rgb = DARK_TEAL
        p_pr.space_before = Pt(8)

        for feat in features:
            p = tf.add_paragraph()
            p.text = feat
            p.font.name = "Arial"
            p.font.size = Pt(10.5)
            p.font.color.rgb = SLATE
            p.space_before = Pt(6)


# ==========================================
# SLIDE 5: HỆ SINH THÁI MẠNG XÃ HỘI
# ==========================================
def build_slide_5(prs: Presentation):
    slide = add_blank_slide(prs)
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Hệ Sinh Thái Truyền Thông", "Mạng Lưới Kênh Mạng Xã Hội Của TravelFrame", "Nơi du khách quốc tế tìm kiếm cảm hứng, cập nhật cẩm nang du lịch và đặt tour trải nghiệm tại Hà Nội.")

    channels = [
        ("📸 Instagram", "@travelframe.vn", "Kênh lõi phân phối video Reels & nhiếp ảnh nghệ thuật. Hơn 41.900+ views/tháng từ du khách quốc tế.", Inches(0.8), Inches(1.9)),
        ("📘 Facebook Fanpage", "TravelFrame Hanoi", "Cổng kết nối cộng đồng, chia sẻ lịch trình du lịch, đánh giá của du khách và đặt tour trực tiếp.", Inches(4.8), Inches(1.9)),
        ("🎵 TikTok", "@travelframe.vn", "Kênh video ngắn bắt trọn xu hướng ẩm thực ngõ nhỏ, cà phê đường tàu và góc check-in ít người biết.", Inches(8.8), Inches(1.9)),
        ("🧵 Threads", "@travelframe.vn", "Chia sẻ mẹo du lịch thực tế, giải đáp nhanh các thắc mắc về thời tiết, di chuyển, giá cả cho khách.", Inches(0.8), Inches(4.5)),
        ("🌐 Website & SEO Hub", "travelframe.sidevn.com", "Cẩm nang ẩm thực phố cổ, điểm chụp ảnh đẹp nhất, booking configurator và chính sách minh bạch.", Inches(4.8), Inches(4.5)),
        ("🎬 YouTube Shorts", "@TravelFrameVN", "Thước phim 4K chất lượng điện ảnh, giới thiệu vẻ đẹp văn hóa và con người Hà Nội tới bạn bè quốc tế.", Inches(8.8), Inches(4.5)),
    ]

    for icon, handle, desc, left, top in channels:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = BORDER_COLOR

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = Inches(0.25)

        p = tf.paragraphs[0]
        p.text = icon
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = DARK_TEAL

        p_h = tf.add_paragraph()
        p_h.text = handle
        p_h.font.name = "Arial"
        p_h.font.size = Pt(11)
        p_h.font.bold = True
        p_h.font.color.rgb = TEAL
        p_h.space_before = Pt(2)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Arial"
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = SLATE
        p_d.space_before = Pt(6)


# ==========================================
# SLIDE 6: CONTENT WIN LÀM TESTIMONIAL
# ==========================================
def build_slide_6(prs: Presentation):
    slide = add_blank_slide(prs)
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Bằng Chứng Thực Tế & Testimonials", "Content Win Của TravelFrame Là Bảo Chứng Cho Niềm Tin", "Những video hữu ích đạt hàng chục ngàn lượt xem đi kèm phản hồi thực chất từ du khách quốc tế.")

    wins = [
        (
            "🎬 Content Win #1: Video '5 Tourist Scams You Need to Avoid in Hanoi' (14.900 views, 126 shares)",
            "\"Nhờ video này mà tụi mình tránh được taxi dù chặt chém và hàng rong ép giá ở Phố Cổ. Khi book tour cùng TravelFrame, Host Mike đưa đi ăn đúng những quán bún chả, bánh mì địa phương ngon nhất mà giá cực kỳ hạt dẻ!\"",
            "— Sarah & Liam (Du khách từ Melbourne, Úc)",
            Inches(1.9)
        ),
        (
            "🎬 Content Win #2: Video 'Hanoi Train Street 2026 Quick Guide' (4.400 views, Share rate 1,82%)",
            "\"Xem video hướng dẫn giờ tàu chạy của TravelFrame rồi book luôn tour buổi chiều. Được Host dẫn vào đúng quán cà phê view ban công tuyệt đẹp tầng 2, uống cà phê trứng ngắm tàu chạy an toàn mà không bị ai chèo kéo!\"",
            "— Jennifer K. (Solo Traveler từ California, Mỹ)",
            Inches(3.6)
        ),
        (
            "🎬 Content Win #3: Video 'POV: Secret Alleys & Coffee Culture' (7.300 views, 24 bình luận)",
            "\"Host dẫn tụi mình luồn qua những con ngõ nhỏ rêu phong mà du khách bình thường không bao giờ biết. Bộ ảnh và video Reel dọc 4K nhận được ngay sau 24h đẹp như cảnh trong phim điện ảnh!\"",
            "— Alex & Marcus (Travel Creators từ London, Anh)",
            Inches(5.3)
        ),
    ]

    for title, quote, author, top in wins:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top, Inches(11.7), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = TEAL

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = Inches(0.18)

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.name = "Arial"
        p1.font.size = Pt(11.5)
        p1.font.bold = True
        p1.font.color.rgb = DARK_TEAL

        p2 = tf.add_paragraph()
        p2.text = quote
        p2.font.name = "Arial"
        p2.font.size = Pt(10.5)
        p2.font.italic = True
        p2.font.color.rgb = SLATE
        p2.space_before = Pt(3)

        p3 = tf.add_paragraph()
        p3.text = author
        p3.font.name = "Arial"
        p3.font.size = Pt(10)
        p3.font.bold = True
        p3.font.color.rgb = TEAL
        p3.space_before = Pt(3)


# ==========================================
# SLIDE 7: CAM KẾT SLA & ĐẶT TOUR 3 BƯỚC
# ==========================================
def build_slide_7(prs: Presentation):
    slide = add_blank_slide(prs)
    set_slide_background(slide, NAVY)

    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
    tf = tb.text_frame
    p0 = tf.paragraphs[0]
    p0.text = "QUY TRÌNH ĐẶT TOUR & CAM KẾT CHẤT LƯỢNG"
    p0.font.name = "Arial"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = RGBColor(56, 189, 248)

    p1 = tf.add_paragraph()
    p1.text = "Ba Bước Tinh Gọn Để Lưu Giữ Kỷ Niệm Cùng TravelFrame"
    p1.font.name = "Arial"
    p1.font.size = Pt(23)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_before = Pt(3)

    steps = [
        ("Bước 1: Chọn Phong Cách", "Lựa chọn ngày đi, số lượng khách và 1-2 Moods yêu thích (Heritage, Everyday, In Style, After Dark).", Inches(0.8)),
        ("Bước 2: Tinh Chỉnh Cùng Host", "Nhắn tin qua WhatsApp hoặc Instagram DM để xác nhận địa điểm đón tại khách sạn, khung giờ và yêu cầu riêng.", Inches(4.8)),
        ("Bước 3: Tận Hưởng & Nhận Ảnh", "Gặp gỡ Host tại khách sạn. Nhận 3-5 ảnh sneak peek trong 24h và trọn bộ gallery bảo mật trong 48h.", Inches(8.8)),
    ]

    for title, desc, left in steps:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.9), Inches(3.7), Inches(2.3))
        card.fill.solid()
        card.fill.fore_color.rgb = SLATE
        card.line.color.rgb = TEAL
        tf_s = card.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_right = tf_s.margin_top = Inches(0.3)

        p = tf_s.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = RGBColor(56, 189, 248)

        p_d = tf_s.add_paragraph()
        p_d.text = desc
        p_d.font.name = "Arial"
        p_d.font.size = Pt(10.5)
        p_d.font.color.rgb = RGBColor(226, 232, 240)
        p_d.space_before = Pt(8)

    cta = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.6), Inches(11.7), Inches(2.3))
    cta.fill.solid()
    cta.fill.fore_color.rgb = DARK_TEAL
    cta.line.color.rgb = RGBColor(56, 189, 248)
    tf_c = cta.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = Inches(0.4)

    p_c1 = tf_c.paragraphs[0]
    p_c1.text = "Sẵn Sàng Biến Chuyến Đi Hà Nội Của Bạn Thành Thước Phim Điện Ảnh?"
    p_c1.font.name = "Arial"
    p_c1.font.size = Pt(18)
    p_c1.font.bold = True
    p_c1.font.color.rgb = WHITE

    p_c2 = tf_c.add_paragraph()
    p_c2.text = "Mỗi Host chỉ nhận tối đa 2 tour/ngày để đảm bảo chất lượng và sự chu đáo tối đa. Miễn phí hủy tour trước 48 giờ."
    p_c2.font.name = "Arial"
    p_c2.font.size = Pt(11)
    p_c2.font.color.rgb = RGBColor(226, 232, 240)
    p_c2.space_before = Pt(4)

    p_c3 = tf_c.add_paragraph()
    p_c3.text = "💬 WhatsApp Concierge: +84 (0) 904 888 644   •   📸 Instagram: @travelframe.vn   •   🌐 Website: travelframe.sidevn.com"
    p_c3.font.name = "Arial"
    p_c3.font.size = Pt(12)
    p_c3.font.bold = True
    p_c3.font.color.rgb = RGBColor(255, 224, 80)
    p_c3.space_before = Pt(12)


def main():
    prs = create_deck()
    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    build_slide_7(prs)
    prs.save(str(OUTPUT))
    print(f"Successfully generated Vietnamese Customer Pitching Catalogue PPTX at: {OUTPUT}")


if __name__ == "__main__":
    main()
